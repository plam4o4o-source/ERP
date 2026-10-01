# -*- coding: utf-8 -*-
"""Общите помощни функции на трите Excel импорта: палетни карти от справка
за поръчки (routes_pallet_extra), редове за фактура (routes_invoices) и
справочника материали (materials/routes_materials).

Одит (01.10.2026, Q5): до тази версия същите функции бяха копирани дословно
в два-три модула и вече се бяха разминали по коментари. Поведението и
текстовете пред оператора са непроменени — един и същ оператор качва и
трите файла и трябва да получава едни и същи съобщения.

Лимитите (HEADER_SCAN_ROWS/MAX_IMPORT_DATA_ROWS) се подават като параметри,
защото всеки импорт пази собствено копие на модулно ниво (тестовете ги
подменят поотделно)."""
import io
import itertools
import zipfile

from flask_babel import gettext as _

from appcore import ensure_xlsx_within_limits

#: Заглавният ред се търси най-много в толкова реда от началото на листа.
HEADER_SCAN_ROWS = 10
#: Таван на редовете данни на един импорт (пази паметта на процеса).
MAX_IMPORT_DATA_ROWS = 5000
#: Колко реда след тавана се преглеждат, за да се реши дали файлът наистина
#: има още ДАННИ (празни, но форматирани редове не са данни).
TRUNCATION_SCAN_ROWS = 100000

# Заглавията на колоните в справка за поръчки (палетни карти и фактури) —
# сравнение след смъкване до малки букви.
ORDER_HEADERS = ("order no", "order number", "orderno")
POS_HEADERS = ("pos", "position")
# Одит (01.10.2026, U10): „Material“ и сродните заглавия се приемат за код.
REFERENCE_HEADERS = ("reference", "material", "material code", "material no",
                     "material number", "part id", "abb part id")
QTY_HEADERS = ("open qty", "qty", "quantity")


def open_workbook(file_bytes):
    """Проверява разархивирания размер (XlsxTooLargeError) и отваря книгата
    в read_only/data_only режим. Всяка друга грешка на openpyxl се пропуска
    към извикващия — той решава как да я съобщи."""
    from openpyxl import load_workbook

    ensure_xlsx_within_limits(file_bytes)
    return load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)


def cellstr(v):
    """Клетка към низ, без излишно „.0“ за цели числа, записани като float.

    Дробният float се закръгля до 6 знака (2.9000000000000004 → „2.9“), но
    ако това би превърнало ненулева стойност в „0“ (8.7e-09), се връща
    суровият запис — редът остава разпознат като неразчитаемо число."""
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        text = ("%.6f" % v).rstrip("0").rstrip(".")
        if text in ("", "0", "-0") and v != 0:
            return str(v).strip()
        return text or "0"
    return str(v).strip()


def norm_header(v):
    """Заглавие на колона към сравним вид: малки букви, без нови редове и
    без повтарящи се празни места."""
    return " ".join(cellstr(v).lower().split())


def reset_sheet_dimensions(ws):
    """openpyxl в `read_only=True` вярва на записания във файла
    `<dimension ref=…>` и не чете извън него — някои генератори записват
    грешен размер и листът тихо се орязва. Обикновен лист няма този метод."""
    reset = getattr(ws, "reset_dimensions", None)
    if reset is not None:
        reset()


def row_has_data(row):
    """Дали редът има поне една непразна клетка (след трим)."""
    return any(cellstr(c) != "" for c in (row or ()))


def more_data_follows(rows, scan_rows=None):
    """Дали в итератора `rows` има ред с ДАННИ (не просто празен форматиран
    ред). Прегледът е ограничен до `scan_rows` реда; ако и след тях има още
    редове, по-безопасно е да предупредим за орязване."""
    scan_rows = TRUNCATION_SCAN_ROWS if scan_rows is None else scan_rows
    for row in itertools.islice(rows, scan_rows):
        if row_has_data(row):
            return True
    return next(rows, None) is not None


def read_limited_rows(ws, header_scan_rows=HEADER_SCAN_ROWS,
                      max_rows=MAX_IMPORT_DATA_ROWS):
    """Чете само толкова реда, колкото могат да бъдат използвани (заглавие +
    таван на данните + един за откриване на орязване), вместо целия лист.

    Връща (rows, exhausted): `exhausted` е True, когато след прочетеното
    няма повече редове с данни (тоест НЯМА орязване)."""
    reset_sheet_dimensions(ws)
    it = ws.iter_rows(values_only=True)
    limit = header_scan_rows + max_rows + 1
    rows = list(itertools.islice(it, limit))
    exhausted = not more_data_follows(it) if len(rows) == limit else True
    return rows, exhausted


def find_col(header_lower, *names):
    """Индексът на първата колона, чието заглавие съвпада с някое от
    `names` (по реда на `names`), или None."""
    for name in names:
        for i, h in enumerate(header_lower):
            if h == name:
                return i
    return None


def locate_header(rows, required, header_scan_rows=HEADER_SCAN_ROWS):
    """Търси заглавния ред сред първите `header_scan_rows` реда: първият,
    в който има колона от ВСЯКА група в `required` (кортеж от кортежи с
    имена). Връща (индекс, суровите заглавия като низове); без намерен ред
    — (0, заглавията на първия ред)."""
    for idx in range(min(header_scan_rows, len(rows))):
        candidate = [cellstr(c) for c in (rows[idx] or [])]
        lower = [h.lower() for h in candidate]
        if all(find_col(lower, *names) is not None for names in required):
            return idx, candidate
    return 0, [cellstr(c) for c in (rows[0] or [])] if rows else []


def header_found_warning(header_row):
    """Предупреждение „заглавието не е на първия ред“; `header_row` е
    1-базиран номер на реда."""
    return (_("Заглавният ред е открит на ред %d от файла (пропуснати са "
              "%d реда над него) — проверете дали разпознатите данни са "
              "правилни.") % (header_row, header_row - 1))


def missing_code_column_warning(header):
    """Одит (01.10.2026, U10): без разпозната колона с код редовете се
    зареждаха с празен код („тегло намерено за 0“), без да се каже защо."""
    found = ", ".join("„%s“" % h for h in header if h) or "—"
    return _("Колоната с код на материала („Reference“ или „Material“) не е "
             "намерена; открити колони: %(found)s. Редовете са заредени без код.") % {
                 "found": found}


def merged_cells_warning():
    """Краткото предупреждение за обединени клетки (фактури/справочник)."""
    return _("Файлът съдържа обединени клетки — стойности извън първата "
             "клетка на обединен диапазон може да липсват.")


def _sheets_contain(file_bytes, needles):
    """Дали суровият XML на някой лист съдържа някой от `needles`. Чете се
    направо архивът — openpyxl в read_only режим не излага merged_cells, а
    втори пълен прочит само за формулите би върнал спестената памет.
    Повреден архив → False (load_workbook дава по-конкретна грешка)."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            for name in zf.namelist():
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                    xml = zf.read(name)
                    if any(n in xml for n in needles):
                        return True
    except Exception:
        return False
    return False


def has_formulas(file_bytes):
    """Формулна клетка без кеширана стойност се чете като None при
    `data_only=True` — файл, генериран извън Excel, се внася с празни
    стойности. Елементът `<f>` в `<c>` е самата формула."""
    return _sheets_contain(file_bytes, (b"<f>", b"<f "))


def has_merged_cells(file_bytes):
    """Обединена клетка има стойност само в горния ляв ъгъл на диапазона —
    останалите клетки се четат като None."""
    return _sheets_contain(file_bytes, (b"<mergeCell ",))


def header_row_hint(header_scan_rows=HEADER_SCAN_ROWS):
    """Допълнение към „няма разпознаваеми колони“: къде се търси заглавието."""
    return _("Заглавният ред трябва да е в първите %d реда на листа.") % header_scan_rows


def formula_hint():
    """Защо иначе валиден файл се внася с празни стойности (has_formulas)."""
    return _("Файлът съдържа формули без запазени стойности — отворете го и го "
             "запишете от Excel, след което опитайте отново.")


def orders_not_recognized_error(file_bytes, header_scan_rows=HEADER_SCAN_ROWS):
    """Грешката при справка за поръчки без разпознаваеми колони — с
    подсказката къде се търси заглавието и, ако е причината, за формулите."""
    msg = _("Файлът не съдържа разпознаваеми колони (Order No, Pos, Reference, "
            "Reference Desc, Open Qty) или редове за импорт.") + " " + header_row_hint(header_scan_rows)
    if has_formulas(file_bytes):
        msg += " " + formula_hint()
    return msg


def parse_sheets(wb, parser):
    """Обхожда листовете подред и връща резултата от ПЪРВИЯ с разпознати
    колони: (parsed, warnings, sheet_title); при нито един — (None, [], None).
    Предупрежденията на неразпознатите (декоративни) листове не се показват."""
    sheets = list(wb.worksheets)
    for ws in sheets:
        parsed, warnings = parser(ws)
        if parsed:
            if len(sheets) > 1:
                warnings.insert(0, _("Данните са прочетени от лист „%(sheet)s“ "
                                     "(файлът съдържа %(count)d листа).")
                                % {"sheet": ws.title, "count": len(sheets)})
            return parsed, warnings, ws.title
    return None, [], None
