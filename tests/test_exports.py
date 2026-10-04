# -*- coding: utf-8 -*-
"""Одит (04.10.2026, група EXPORT): PDF/Excel износ на документите.

X1 — колоните на PDF таблицата не се застъпват и не излизат от рамката,
     заглавията не се режат, думите не се чупят по средата;
X2 — фактурата за Бразилия се изнася БЕЗ „Общо тегло“ (както бланката);
X3 — Excel: числа с разделител за хиляди и без „10.“, суми в евро и дати
     като истински стойности;
X4 — Excel: настройки за печат, повтарящ се заглавен ред, рамки, пренос;
X5 — PDF: празните полета не се печатат, новите редове се пазят, редът
     ОБЩО/TOTAL на опаковъчния лист и палетната карта;
R6 — число с над 308 цифри не изчезва от Excel;
F10 — името на изнесения файл на палетна карта не зависи от адресната книга;
I5 — етикетите на полетата в предупрежденията са на езика на интерфейса;
UX-№4 — Excel износ на филтрирания списък с документи.
"""
import io
import json
import re
from datetime import date, datetime

import pypdf
import pytest
from openpyxl import load_workbook

from conftest import post_with_csrf

# A4 (595.28 pt) минус страничните полета на @page (2 × 1.2 cm).
_FRAME_LEFT = 1.2 * 72 / 2.54
_FRAME_RIGHT = 595.28 - _FRAME_LEFT

NO_HEADER = {"consignee_name": "Embraer S.A.", "doc_date": "2026-10-04",
             "consignee_address": "Av. Brigadeiro Faria Lima, 2170\nSão José dos Campos"}


def _issue(client, url, fields):
    resp = post_with_csrf(client, url, fields, csrf_source_url=url, follow_redirects=False)
    assert resp.status_code == 302, resp.data[:400]
    m = re.search(r"/doc/(\d+)$", resp.headers["Location"])
    assert m, resp.headers["Location"]
    return int(m.group(1))


def _norway_items(n=6):
    return [{"hs_code": "8483109590",
             "description": "Задвижващ вал за редуктор, стомана 42CrMo4, тип %d" % i,
             "pallet_no": str(i), "po_no": "45000000%02dA2" % (i * 7), "pos": str(i * 10),
             "material_code": "C0000%06d-%02d" % (12345 * i, i),
             "qty": str(25 + 25 * i), "unit_price": "%d.%02d" % (100 + 37 * i, i)}
            for i in range(1, n + 1)]


def _spans(pdf_bytes):
    """[(страница, y, x, текст, ширина)] за всеки изрисуван низ."""
    from reportlab.pdfbase import pdfmetrics

    import pdf_export
    out = []
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    for index, page in enumerate(reader.pages):
        def visit(text, cm, tm, font, size, index=index):
            if not text.strip():
                return
            x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
            y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
            base = (font or {}).get("/BaseFont", "") if hasattr(font, "get") else ""
            bold = "Bold" in str(base)
            name = pdf_export._metric_font(bold)
            eff = size * (tm[0] or 1.0)
            out.append((index, round(y, 1), x, text.strip(),
                        pdfmetrics.stringWidth(text.strip(), name, eff)))
        page.extract_text(visitor_text=visit)
    return out


# ---------------------------------------------------------------- X1

def test_x1_pdf_columns_do_not_overlap_and_stay_inside_the_frame(admin_client):
    """Фактура за Норвегия с реални стойности: преди поправката P.O NO и кодът
    на материала се застъпваха („4500000007A2C0000…“), HS кодът започваше
    вляво от рамката, заглавията се режеха („Колич“), а думите се чупеха
    по средата („мате/риала“)."""
    doc_id = _issue(admin_client, "/invoice-no/new", dict(
        NO_HEADER, invoice_number="0000012957",
        items_json=json.dumps(_norway_items(), ensure_ascii=False)))
    resp = admin_client.get("/doc/%d/export.pdf" % doc_id)
    assert resp.status_code == 200
    spans = [s for s in _spans(resp.data) if s[1] > 60]  # без колонтитула
    texts = [s[3] for s in spans]
    for whole in ("4500000007A2", "C0000012345-01", "8483109590", "Количество",
                  "материала"):
        assert whole in texts, "%r не е изрисувано цяло: %r" % (whole, texts[:80])
    assert not any(t in ("мате", "риала", "Колич") for t in texts)
    for page, y, x, text, width in spans:
        assert x >= _FRAME_LEFT - 1, "%r започва вляво от рамката (x=%.1f)" % (text, x)
        assert x + width <= _FRAME_RIGHT + 1, "%r излиза вдясно от рамката" % text
    # Низовете на един и същ ред (обща базова линия) не се застъпват.
    by_line = {}
    for page, y, x, text, width in spans:
        by_line.setdefault((page, y), []).append((x, width, text))
    for line in by_line.values():
        line.sort()
        for (x1, w1, t1), (x2, _w2, t2) in zip(line, line[1:]):
            assert x1 + w1 <= x2 + 0.5, "%r се застъпва с %r" % (t1, t2)


def test_x1_a_token_wider_than_its_column_is_wrapped_by_measure():
    """Дума, по-широка от колоната, се реже на части, всяка от които се
    побира — без загуба на знаци; думите, които се побират, остават цели."""
    import pdf_export
    token = "MATERIALCODE-ÄÖÜß-ĞŞİÇ-" + "X" * 40
    width = 80.0
    lines = pdf_export.wrap_cell_lines("кратко " + token, width, size=9.0)
    inner = width - pdf_export._PDF_CELL_EXTRA_PT
    assert lines[0] == "кратко"
    assert "".join(lines[1:]) == token
    for line in lines:
        assert pdf_export._text_width(line, False, 9.0) <= inner + 0.5
    assert pdf_export.wrap_cell_lines("ред 1\nред 2", 200) == ["ред 1", "ред 2"]


def test_x1_no_cjk_mid_word_wrapping_in_the_template():
    import os
    from conftest import ROOT
    html = open(os.path.join(ROOT, "templates", "pdf_export.html"), encoding="utf-8").read()
    css = html[html.index("<style>"):html.index("</style>")]
    css = re.sub(r"\{#.*?#\}", "", css, flags=re.S)
    assert "-pdf-word-wrap" not in css


def test_x1_widths_follow_the_content_and_fill_the_frame():
    import pdf_export
    cols = [("hs_code", "HS code"), ("description", "Описание на материала"),
            ("po_no", "P.O NO"), ("qty", "Количество")]
    items = [{"hs_code": "8483109590", "description": "дълго описание " * 8,
              "po_no": "4500000007A2", "qty": "50"}]
    widths, size = pdf_export.pdf_table_plan(cols, items)
    assert abs(sum(widths) - pdf_export._PDF_FRAME_WIDTH_PT) < 0.5
    for (key, label), w in zip(cols, widths):
        value = items[0][key] if key != "description" else "описание"
        need = max(pdf_export._text_width(value, False, size),
                   max(pdf_export._text_width(word, True, size) for word in label.split()))
        assert w - pdf_export._PDF_CELL_EXTRA_PT >= need - 0.5, (key, w, need)
    assert widths[1] == max(widths), "свободният текст поема остатъка"


# ---------------------------------------------------------------- X2

def test_x2_brazil_exports_have_no_total_weight(admin_client):
    items = [{"hs_code": "7326909890", "po_no": "4500000007", "pos": "10",
              "net_weight": "1.25", "material_code": "MAT-1", "qty": "4", "unit_price": "2.5"}]
    doc_id = _issue(admin_client, "/invoice-br/new", dict(
        NO_HEADER, invoice_number="BR-1", items_json=json.dumps(items)))
    wb = load_workbook(io.BytesIO(admin_client.get("/doc/%d/export.xlsx" % doc_id).data))
    values = [c.value for row in wb.active.iter_rows() for c in row if c.value is not None]
    assert "Общо тегло, кг" not in values
    assert 5.0 not in values, "общото тегло (4 × 1.25) не бива да е в износа"
    assert "Нето тегло, кг/бр" in values  # колоната от бланката остава
    pdf = admin_client.get("/doc/%d/export.pdf" % doc_id).data
    text = "".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf)).pages)
    assert "Общо" not in text.replace("Обща цена", "")


# ---------------------------------------------------------------- X3 / X4

def _sheet(client, doc_id):
    resp = client.get("/doc/%d/export.xlsx" % doc_id)
    assert resp.status_code == 200
    return load_workbook(io.BytesIO(resp.data)).active


def _row_of(ws, first_value):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == first_value:
            return row[0].row
    raise AssertionError("няма ред, започващ с %r" % first_value)


def test_x3_numbers_dates_and_euro_totals_are_real_values(admin_client):
    items = [{"hs_code": "7326", "po_no": "PO1", "pos": "10", "material_code": "M1",
              "qty": "1200", "unit_price": "1234.5"},
             {"hs_code": "7326", "po_no": "PO1", "pos": "20", "material_code": "M2",
              "qty": "2.5", "unit_price": "0.0125"}]
    doc_id = _issue(admin_client, "/invoice-dubai/new", dict(
        NO_HEADER, invoice_number="DX-1", items_json=json.dumps(items)))
    ws = _sheet(admin_client, doc_id)
    date_cell = ws.cell(row=_row_of(ws, "Дата"), column=2)
    assert isinstance(date_cell.value, (date, datetime))
    assert date_cell.value.strftime("%Y-%m-%d") == "2026-10-04"
    assert date_cell.number_format == "dd.mm.yyyy"
    first = _row_of(ws, "HS code") + 1
    qty = ws.cell(row=first, column=5)
    assert qty.value == 1200 and qty.number_format == "#,##0", qty.number_format
    assert ws.cell(row=first + 1, column=5).number_format == "#,##0.0"
    assert ws.cell(row=first, column=6).number_format == "#,##0.00###"
    total = _row_of(ws, "TOTAL")
    total_price = ws.cell(row=total, column=7)
    assert isinstance(total_price.value, float), "общата сума е текст: %r" % total_price.value
    assert abs(total_price.value - (1200 * 1234.5 + 2.5 * 0.0125)) < 0.01
    assert total_price.number_format == '#,##0.00 "€"'


def test_x3_waybill_freight_is_a_number_in_euro(admin_client):
    doc_id = _issue(admin_client, "/waybill/new", {
        "consignee_name": "Строймат АД", "transport_price": "1450.50", "extra_costs": "35",
        "established_date": "2026-10-04",
        "items_json": json.dumps([{"description": "Профили", "qty": "3", "weight": "100.5"}])})
    ws = _sheet(admin_client, doc_id)
    cell = ws.cell(row=_row_of(ws, "Превозна цена, EUR"), column=2)
    assert cell.value == 1450.5 and cell.number_format == '#,##0.00 "€"'


def test_x4_print_setup_header_repeat_borders_and_wrap(admin_client):
    doc_id = _issue(admin_client, "/invoice-no/new", dict(
        NO_HEADER, invoice_number="NO-1",
        items_json=json.dumps(_norway_items(3), ensure_ascii=False)))
    ws = _sheet(admin_client, doc_id)
    header = _row_of(ws, "HS code")
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.print_title_rows == "$%d:$%d" % (header, header)
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A%d:" % header)
    desc = ws.cell(row=header + 1, column=2)
    assert desc.alignment.wrap_text
    assert desc.border.left.style == "thin" and desc.border.bottom.style == "thin"
    addr = ws.cell(row=_row_of(ws, "Адрес получател"), column=2)
    assert addr.alignment.wrap_text and addr.border.top.style == "thin"


# ---------------------------------------------------------------- X5

def _pdf_text(client, doc_id):
    data = client.get("/doc/%d/export.pdf" % doc_id).data
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(data)).pages)


def test_x5_pdf_skips_empty_fields_and_keeps_newlines(admin_client):
    doc_id = _issue(admin_client, "/packing/new", {
        "receiver_name": "Müller Logistik GmbH",
        "sender_address": "12 Industrialna Str.\n5300 Gabrovo, Bulgaria",
        "items_json": json.dumps([{"description": "Скоби", "qty": "1"}])})
    text = _pdf_text(admin_client, doc_id)
    assert "Имейл получател" not in text, "празното поле е отпечатано"
    lines = [ln.strip() for ln in text.splitlines()]
    assert "12 Industrialna Str." in lines and "5300 Gabrovo, Bulgaria" in lines


def test_x5_packing_and_pallet_get_a_totals_row(admin_client):
    packing = _issue(admin_client, "/packing/new", {
        "receiver_name": "R", "total_packages": "2", "total_volume": "1.5",
        "total_net": "300", "total_gross": "321.5",
        "items_json": json.dumps([{"description": "A", "qty": "1", "net": "100", "gross": "110"},
                                  {"description": "B", "qty": "2", "net": "200", "gross": "211.5"}])})
    text = _pdf_text(admin_client, packing)
    assert "ОБЩО / TOTAL" in text.replace("\n", " ") and "321.5" in text
    ws = _sheet(admin_client, packing)
    total = _row_of(ws, "ОБЩО / TOTAL")
    header = _row_of(ws, "Вид опаковка")
    gross_col = [ws.cell(row=header, column=c).value for c in range(1, 10)].index("Бруто, кг") + 1
    assert ws.cell(row=total, column=gross_col).value == 321.5
    assert ws.cell(row=total, column=2).value == "2 колета/packages"

    pallet = _issue(admin_client, "/pallet/new", {
        "client_name": "Клиент", "items_json": json.dumps(
            [{"code": "A", "qty": "5"}, {"code": "B", "qty": "7"}])})
    ws = _sheet(admin_client, pallet)
    total = _row_of(ws, "ОБЩО / TOTAL")
    assert ws.cell(row=total, column=3).value == 12
    assert "ОБЩО / TOTAL" in _pdf_text(admin_client, pallet).replace("\n", " ")


# ---------------------------------------------------------------- R6

def test_r6_huge_number_stays_visible_in_excel(admin_client):
    huge = "1" * 400
    doc_id = _issue(admin_client, "/packing/new", {
        "receiver_name": "R", "total_net": huge,
        "items_json": json.dumps([{"description": "A", "qty": huge}])})
    ws = _sheet(admin_client, doc_id)
    assert ws.cell(row=_row_of(ws, "Общо нето, кг"), column=2).value == huge
    first = _row_of(ws, "Вид опаковка") + 1
    assert ws.cell(row=first, column=3).value == huge, "числото изчезна (inf → празна клетка)"


# ---------------------------------------------------------------- F10

def test_f10_pallet_file_name_survives_deleting_the_client(admin_client, db_module):
    con = db_module.get_db()
    con.execute("INSERT INTO clients (name, alias) VALUES (?, ?)", ("Дипласт ООД", "DSP"))
    con.commit()
    con.close()
    doc_id = _issue(admin_client, "/pallet/new", {
        "client_name": "Дипласт ООД", "items_json": json.dumps([{"code": "A", "qty": "1"}])})
    first = admin_client.get("/doc/%d/export.pdf" % doc_id)
    assert "DSP_0001-" in first.headers["Content-Disposition"]
    con = db_module.get_db()
    con.execute("DELETE FROM clients")
    con.commit()
    con.close()
    for ext in ("pdf", "xlsx"):
        resp = admin_client.get("/doc/%d/export.%s" % (doc_id, ext))
        assert "DSP_0001-" in resp.headers["Content-Disposition"], resp.headers["Content-Disposition"]


def test_f10_legacy_pallet_keeps_the_alias_after_its_first_export(admin_client, db_module):
    """Документ отпреди поправката (без записан псевдоним) получава
    псевдонима при първия износ и го пази и след изтриване на клиента."""
    con = db_module.get_db()
    con.execute("INSERT INTO clients (name, alias) VALUES (?, ?)", ("Стар клиент", "OLD"))
    con.execute("INSERT INTO documents (doc_type, number, year, seq, barcode, public_token,"
                " data, created_by) VALUES ('pallet', '0009/2026', 2026, 9, 'B9', 'T9', ?, 1)",
                (json.dumps({"client_name": "Стар клиент", "items": []}, ensure_ascii=False),))
    con.commit()
    doc_id = con.execute("SELECT id FROM documents WHERE barcode = 'B9'").fetchone()[0]
    con.close()
    assert "OLD_0009-2026" in admin_client.get(
        "/doc/%d/export.xlsx" % doc_id).headers["Content-Disposition"]
    con = db_module.get_db()
    con.execute("DELETE FROM clients")
    con.commit()
    con.close()
    assert "OLD_0009-2026" in admin_client.get(
        "/doc/%d/export.xlsx" % doc_id).headers["Content-Disposition"]


# ---------------------------------------------------------------- I5

def test_i5_field_labels_in_warnings_follow_the_ui_language(admin_client):
    with admin_client.session_transaction() as sess:
        sess["lang"] = "en"
    resp = post_with_csrf(admin_client, "/cmr/new", {
        "consignee_name": "R", "weight": "1.234,56"},
        csrf_source_url="/cmr/new", follow_redirects=True)
    body = resp.get_data(as_text=True)
    # Етикетът в предупреждението стои в кавички „…“ (бланката под него е
    # двуезична по замисъл и съдържа и двата текста).
    assert "„Gross weight, kg“" in body
    assert "„Бруто тегло, кг“" not in body


def test_i5_conflict_dialog_labels_are_translated(flask_app):
    import routes_documents as rd
    with flask_app.test_request_context():
        from flask import session
        session["lang"] = "en"
        lines = rd._conflict_differences("cmr", {"weight": "1"}, {"weight": "2"})
    assert lines and lines[0].startswith("Gross weight, kg"), lines


def test_i5_field_labels_are_marked_for_extraction():
    import os
    from conftest import ROOT
    src = open(os.path.join(ROOT, "routes_documents.py"), encoding="utf-8").read()
    assert '(N_("Бруто тегло, кг"), "weight")' in src
    assert '(N_("Дата на съставяне"), "established_date")' in src


# ---------------------------------------------------------------- UX-№4

def test_list_export_follows_the_filters_and_has_totals(admin_client, db_module):
    cmr1 = _issue(admin_client, "/cmr/new", {"consignee_name": "Алфа", "weight": "1200.5",
                                             "packages": "10", "volume": "3.25",
                                             "established_date": "2026-10-01"})
    _issue(admin_client, "/cmr/new", {"consignee_name": "Бета", "weight": "800",
                                      "packages": "5"})
    _issue(admin_client, "/packing/new", {"receiver_name": "Алфа", "total_gross": "50",
                                          "order_no": "PO-77", "invoice_no": "INV-1",
                                          "items_json": json.dumps([{"description": "x"}])})
    con = db_module.get_db()
    con.execute("INSERT INTO document_attachments (document_id, token, filename, ext, size)"
                " VALUES (?, 'tok-a', 'cmr.pdf', 'pdf', 10)", (cmr1,))
    con.commit()
    con.close()

    resp = admin_client.get("/docs/export.xlsx?type=cmr")
    assert resp.status_code == 200
    assert resp.mimetype.endswith("spreadsheetml.sheet")
    ws = load_workbook(io.BytesIO(resp.data)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0][:4] == ["Дата", "Тип", "№", "Клиент"]
    body = rows[1:-1]
    assert len(body) == 2, "филтърът по тип не е приложен: %r" % body
    by_client = {r[3]: r for r in body}
    assert by_client["Алфа"][11] == "Да" and by_client["Бета"][11] == "Не"
    assert by_client["Алфа"][0].strftime("%d.%m.%Y") == "01.10.2026"
    total = rows[-1]
    assert total[0] == "Общо" and total[2] == 2
    assert total[6] == 15 and abs(total[7] - 2000.5) < 1e-9
    assert ws.freeze_panes == "A2"

    ws = load_workbook(io.BytesIO(admin_client.get("/docs/export.xlsx?q=PO-77").data)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert len(rows) == 3 and rows[1][4] == "PO-77" and rows[1][5] == "INV-1"


def test_list_export_requires_login_and_excludes_invoices(client, employee_client):
    assert client.get("/docs/export.xlsx").status_code in (302, 401)
    _issue(employee_client, "/invoice-br/new", {
        "consignee_name": "ABB", "invoice_number": "BR-9",
        "items_json": json.dumps([{"material_code": "M", "qty": "1", "unit_price": "1"}])})
    resp = employee_client.get("/docs/export.xlsx")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert ws.max_row == 2, "фактурите не принадлежат на списъка с документи"


def test_list_page_has_the_excel_button_that_submits_the_current_filters(admin_client):
    body = admin_client.get("/docs").get_data(as_text=True)
    m = re.search(r'<button[^>]*formaction="/docs/export\.xlsx"[^>]*>', body)
    assert m and 'formmethod="get"' in m.group(0), "бутонът за Excel износ на списъка липсва"
    from conftest import app_js_source
    assert 'e.submitter.hasAttribute("formaction")' in app_js_source(), (
        "живото търсене би спряло изпращането на формата към износа")


def test_packing_totals_fall_back_to_the_row_sums_like_the_print(admin_client):
    doc_id = _issue(admin_client, "/packing/new", {
        "receiver_name": "R", "total_gross": "999",
        "items_json": json.dumps([{"description": "A", "qty": "2", "net": "1.5", "gross": "2"},
                                  {"description": "B", "qty": "3", "net": "2.25", "gross": "3"}])})
    ws = _sheet(admin_client, doc_id)
    header = _row_of(ws, "Вид опаковка")
    names = [ws.cell(row=header, column=c).value for c in range(1, 10)]
    total = _row_of(ws, "ОБЩО / TOTAL")
    value = {n: ws.cell(row=total, column=i + 1).value for i, n in enumerate(names)}
    assert value["Брой"] == 5, "общото количество е сборът на редовете"
    assert value["Нето, кг"] == 3.75, "празно въведено нето → сборът на редовете"
    assert value["Бруто, кг"] == 999, "въведеното бруто се печата както е въведено"


def test_invoice_material_codes_are_saved_uppercase(admin_client, db_module):
    doc_id = _issue(admin_client, "/invoice-dubai/new", {
        "consignee_name": "ABB", "invoice_number": "UP-1",
        "items_json": json.dumps([{"material_code": " glbk-400a ", "qty": "1", "unit_price": "1"}])})
    con = db_module.get_db()
    data = json.loads(con.execute("SELECT data FROM documents WHERE id = ?", (doc_id,)).fetchone()[0])
    con.close()
    assert data["items"][0]["material_code"] == "GLBK-400A"


def test_dualuse_recipient_name_is_used_for_the_client_folder():
    import client_export
    assert client_export.resolve_client_name({"dest_name": "Турски клиент"}) == "Турски клиент"
    assert client_export.resolve_client_name({"consignee_name": "А", "dest_name": "Б"}) == "А"


def test_r6_parser_rejects_non_finite_numbers():
    import appcore
    assert appcore._parse_decimal("1" * 400) is None
    assert appcore._parse_decimal("12,5") == 12.5


@pytest.mark.e2e
def test_list_excel_button_downloads_the_live_filtered_list(page, live_server):
    """UX-№4 в браузъра: живото търсене стеснява списъка без презареждане —
    бутонът „Excel“ трябва да изнесе ТОЧНО показаното (и да не бъде спрян
    от обработчика на живото търсене)."""
    from conftest import e2e_login
    e2e_login(page, live_server)
    for name in ("Алфа ЕООД", "Бета ЕООД"):
        page.goto(live_server + "/cmr/new")
        page.fill("#f-consignee_name", name)
        page.click("#main-doc-form button[type=submit]")
        page.wait_for_url(live_server + "/doc/*")
    page.goto(live_server + "/docs")
    page.fill("#f-q", "Бета")
    page.wait_for_function("() => { const t = document.querySelector('#docs-results').innerText;"
                           " return t.includes('Бета ЕООД') && !t.includes('Алфа ЕООД'); }")
    with page.expect_download() as info:
        page.click("button[formaction$='/docs/export.xlsx']")
    assert info.value.suggested_filename.endswith(".xlsx")
    with open(info.value.path(), "rb") as fh:
        ws = load_workbook(io.BytesIO(fh.read())).active
    clients = [row[3].value for row in ws.iter_rows(min_row=2) if row[0].value != "Общо"]
    assert clients == ["Бета ЕООД"], clients
