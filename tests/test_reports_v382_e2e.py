# -*- coding: utf-8 -*-
"""Одит (15.10.2026, REPORTS): „Справки“ в истински браузър — линкът в
страничната лента, бързите бутони за период, групирането, Excel износът,
линкът на клиента към списъка с документи и картите на телефон (390px)."""
import io
from datetime import date

import pytest

from conftest import e2e_login
from test_reports_v382 import _add

pytestmark = pytest.mark.e2e


def _seed(db_module):
    today = date.today()
    now = today.strftime("%Y-%m-%d 09:00:00")
    _add(db_module, "cmr", {"consignee_name": "Е2Е Клиент ООД", "weight": "1500"}, now, raw=True)
    _add(db_module, "cmr", {"consignee_name": "е2е клиент ООД", "weight": "500"}, now, raw=True)
    _add(db_module, "packing", {"receiver_name": "Е2Е Клиент ООД", "total_gross": "2000"}, now, raw=True)
    for _i in range(3):
        _add(db_module, "pallet", {"client_name": "Е2Е Клиент ООД", "gross": "700"}, now, raw=True)
    _add(db_module, "invoice_dubai", {"consignee_name": "Е2Е Клиент ООД",
                                      "items": [{"qty": "4", "unit_price": "12.5"}]}, now, raw=True)
    _add(db_module, "packing", {"receiver_name": "Само ОЛ ЕАД", "total_gross": "80"}, now, raw=True)
    return today


def test_report_from_the_sidebar_quick_buttons_and_excel(page, live_server, db_module):
    from openpyxl import load_workbook
    today = _seed(db_module)
    e2e_login(page, live_server)
    page.click('#sidebar a.nav-item[href="/reports"]')
    page.wait_for_url(live_server + "/reports")
    assert page.locator("#sidebar a.nav-item.active").get_attribute("href") == "/reports"
    rows = page.locator("table.report-table tr")
    assert rows.count() == 4  # заглавие, двата клиента, „Общо“
    first = page.locator("table.report-table tr").nth(1).locator("td")
    assert first.nth(0).inner_text().strip() == "Е2Е Клиент ООД"
    assert [first.nth(i).inner_text().replace("\xa0", " ").strip() for i in range(1, 6)] == \
        ["2", "3", "2 000", "1", "50.00 €"]
    # „Само ОЛ ЕАД“ няма ЧМР — пратките и брутото са от опаковъчния лист (знак „ОЛ“ с обяснение)
    marks = page.locator("table.report-table tr").nth(2).locator("abbr.report-src")
    assert marks.count() == 2 and marks.first.get_attribute("title") == "по опаковъчните листове"
    assert page.locator("table.report-table tr").nth(1).locator("abbr.report-src").count() == 0
    # „Тази година“ — отбелязан бутон и периодът от януари.
    page.click(".report-quick a:has-text('Тази година')")
    page.wait_for_url("**from=%d-01**" % today.year)
    assert page.locator(".report-quick a.is-active").inner_text().strip() == "Тази година"
    assert page.input_value("#r-from") == "%d-01" % today.year
    # Групиране по месец.
    page.select_option("#r-group", "month")
    page.click(".report-form button[type=submit]:not([formaction])")
    page.wait_for_url("**group=month**")
    assert page.locator("table.report-table th").first.inner_text().strip().lower() == "месец"
    # Excel — същият филтър.
    with page.expect_download() as dl:
        page.click(".report-form button[formaction='/reports/export.xlsx']")
    path = dl.value.path()
    ws = load_workbook(io.BytesIO(open(path, "rb").read())).active
    assert ws.cell(row=1, column=1).value == "Месец"
    total_row = [r for r in ws.iter_rows(values_only=True) if r[0] == "Общо"][0]
    assert list(total_row[1:6]) == [3, 3, 2080, 1, 50.0]
    assert ws.cell(row=1, column=7).value == "Забележка"


def test_client_link_opens_the_filtered_documents(page, live_server, db_module):
    _seed(db_module)
    e2e_login(page, live_server)
    page.goto(live_server + "/reports")
    page.click("table.report-table a:has-text('Е2Е Клиент ООД')")
    page.wait_for_url("**/docs?**")
    assert page.input_value("#f-q") == "Е2Е Клиент ООД"
    # Двете ЧМР (и в двата регистъра) и опаковъчният лист; фактурите — не тук.
    assert page.locator("table.list tr[data-row-id]").count() == 2 + 1 + 3


def test_phone_cards_without_horizontal_scroll_light_and_dark(live_server, db_module, e2e_context_factory):
    _seed(db_module)
    ctx = e2e_context_factory(viewport={"width": 390, "height": 844})
    pg = ctx.new_page()
    e2e_login(pg, live_server)
    for theme in ("light", "dark"):
        pg.goto(live_server + "/reports")
        pg.evaluate("t => document.documentElement.setAttribute('data-theme', t)", theme)
        assert pg.evaluate("document.documentElement.scrollWidth") <= 390
        # заглавният ред е скрит, всеки ред е карта с етикети от data-label
        assert not pg.locator("table.report-table tr").first.is_visible()
        total = pg.locator("table.report-table tr.report-total")
        assert total.is_visible() and total.evaluate("e => getComputedStyle(e).display") == "grid"
        label = pg.locator("table.report-table tr.report-total td[data-label]").first
        assert label.evaluate("e => getComputedStyle(e, '::before').content") != "none"
        # контролите са поне 44px високи (пръст)
        for sel in ("#r-from", "#r-to", "#r-client", "#r-group"):
            assert pg.locator(sel).bounding_box()["height"] >= 43.5, sel
