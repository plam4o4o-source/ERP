# -*- coding: utf-8 -*-
"""v3.82.1, група APPFIX (одит 21.10.2026) — находките от функционалния тест:

APPFIX-1  номер на фактура без регистър/интервали („br-0011“ == „BR-0011“);
APPFIX-2  проверка на тарата в Настройки + нето от справочника > бруто;
APPFIX-3  „Опаковъчен лист от тази поръчка (PO)“ симетрично с интервали в PO;
APPFIX-4  филтър „Клиент“ в Справки — регистър (турски), част от името;
APPFIX-5  филтър `client` в /docs и /invoices (линковете от Справки);
APPFIX-6  „Групирай по клиент“ по ключа на клиента;
APPFIX-7  остарял текст на страницата „Система“;
APPFIX-8  пренасяне на дълъг път (<code>) — виж и _e2e;
APPFIX-9  десетичен знак при автоматично попълване (сървърната част);
APPFIX-10 търсене без диакритика + няколко думи (AND);
APPFIX-11 обърнат период в /docs и /invoices."""
import json
import re
import sqlite3
from urllib.parse import parse_qs, quote, urlsplit

import pytest

from conftest import post_with_csrf, read_source
from test_workflow_shortcuts import _add_material, _issue_pallets, _pull, _rows

ITEM = {"material_code": "MAT-1", "qty": "2", "unit_price": "1.5"}
_SEQ = {"n": 0}


def _post(client, url, fields, follow=False):
    return post_with_csrf(client, url, fields, csrf_source_url=url, follow_redirects=follow)


def _invoice(client, url, number, follow=False, **extra):
    fields = {"consignee_name": "ABB", "invoice_number": number, "items_json": json.dumps([ITEM])}
    fields.update(extra)
    return _post(client, url, fields, follow=follow)


def _numbers(db_module):
    con = db_module.get_db()
    try:
        return sorted(r[0] for r in con.execute(
            "SELECT number FROM documents WHERE doc_type LIKE 'invoice%'"))
    finally:
        con.close()


def _add(db_module, doc_type, data, created_at="2026-07-10 10:00:00", number=None):
    """Документ направо в базата (тригерите попълват client_name и тялото за търсене)."""
    _SEQ["n"] += 1
    n = _SEQ["n"]
    con = db_module.get_db()
    cur = con.execute(
        "INSERT INTO documents (doc_type, number, year, seq, barcode, public_token, data,"
        " created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)",
        (doc_type, number or "A%04d" % n, int(created_at[:4]), n, "APF-%06d" % n, "apf-%d" % n,
         json.dumps(data, ensure_ascii=False), created_at))
    con.commit()
    doc_id = cur.lastrowid
    con.close()
    return doc_id


def _ids(body):
    return {int(x) for x in re.findall(r'data-row-id="(\d+)"', body)}


# ---------------------------------------------------------------- APPFIX-1

def test_invoice_number_of_same_type_differing_only_by_case_is_taken(admin_client, db_module):
    _invoice(admin_client, "/invoice-br/new", "BR-0011")
    for typed in ("br-0011", " Br-0011 "):
        resp = _invoice(admin_client, "/invoice-br/new", typed)
        assert "/invoice-br/new?restore=" in resp.headers["Location"], typed
        body = admin_client.get(resp.headers["Location"]).get_data(as_text=True)
        assert "вече има издаден документ с номер BR-0011" in body
    assert _numbers(db_module) == ["BR-0011"]


def test_invoice_number_of_other_type_differing_by_case_needs_confirmation(admin_client, db_module):
    _invoice(admin_client, "/invoice-no/new", "NO-2026106")
    resp = _invoice(admin_client, "/invoice-dubai/new", "no-2026106")
    assert "restore=" in resp.headers["Location"]
    body = admin_client.get(resp.headers["Location"]).get_data(as_text=True)
    assert "NO-2026106 вече е използван" in body and "Фактура за Норвегия" in body
    assert _numbers(db_module) == ["NO-2026106"]
    # Същият поток за потвърждение — второто „Издай“ със същия номер.
    resp = _invoice(admin_client, "/invoice-dubai/new", "no-2026106")
    assert re.search(r"/doc/\d+$", resp.headers["Location"])
    assert _numbers(db_module) == ["NO-2026106", "no-2026106"]


def test_invoice_number_case_check_on_edit_and_from_preview(admin_client, db_module):
    _invoice(admin_client, "/invoice-br/new", "BR-7")
    resp = _invoice(admin_client, "/invoice-br/new", "BR-8")
    doc_id = int(resp.headers["Location"].rsplit("/", 1)[-1])
    resp = post_with_csrf(admin_client, "/doc/%d/edit" % doc_id, {
        "consignee_name": "ABB", "invoice_number": "br-7", "items_json": json.dumps([ITEM])})
    assert "restore=" in resp.headers["Location"]
    assert _numbers(db_module) == ["BR-7", "BR-8"]
    # Само регистърът на СОБСТВЕНИЯ номер — не е конфликт.
    resp = post_with_csrf(admin_client, "/doc/%d/edit" % doc_id, {
        "consignee_name": "ABB", "invoice_number": "br-8", "items_json": json.dumps([ITEM])})
    assert resp.headers["Location"].endswith("/doc/%d" % doc_id)
    # Предварителен преглед → „Издай“ минава през същата проверка.
    resp = _post(admin_client, "/invoice-br/preview", {
        "consignee_name": "ABB", "invoice_number": " bR-7", "items_json": json.dumps([ITEM])})
    token = resp.headers["Location"].rsplit("/", 1)[-1]
    resp = post_with_csrf(admin_client, "/preview/%s/issue" % token, {},
                          csrf_source_url="/invoice-br/new")
    assert "/invoice-br/new?restore=" in resp.headers["Location"]
    assert _numbers(db_module) == ["BR-7", "br-8"]


def test_invoice_number_check_uses_the_type_year_number_index(db_module):
    src = read_source("routes_documents.py")
    assert "upper(trim(number)) = upper(?)" in src
    con = db_module.get_db()
    plan = " ".join(r[3] for r in con.execute(
        "EXPLAIN QUERY PLAN SELECT doc_type, number FROM documents WHERE year = ?"
        " AND upper(trim(number)) = upper(?) AND doc_type IN (?, ?, ?)",
        (2026, "x", "invoice_br", "invoice_no", "invoice_dubai")))
    con.close()
    assert "idx_documents_type_year_number" in plan, plan


# ---------------------------------------------------------------- APPFIX-2

@pytest.mark.parametrize("raw, stored", [("20", "20"), ("20 кг", "20"), ("25,5 kg", "25,5"),
                                         (" 1 2,5 KG. ", "12,5"), ("0", "0"), ("", ""),
                                         ("500", "500")])
def test_tare_normalization_accepts(db_module, raw, stored):
    assert db_module.normalize_pallet_tare(raw) == stored


@pytest.mark.parametrize("raw", ["abc", "-5", "1000", "500,5", "20 кг 5", "1.234,5", "1e3"])
def test_tare_normalization_rejects(db_module, raw):
    assert db_module.normalize_pallet_tare(raw) is None


def _settings_form(client, **changes):
    import html
    body = client.get("/settings").get_data(as_text=True)
    form = {k: html.unescape(v) for k, v in
            re.findall(r'<input[^>]*name="([a-z_]+)"[^>]*value="([^"]*)"', body)}
    form.update(changes)
    return form


def _stored_tare(db_module):
    con = db_module.get_db()
    try:
        return db_module.get_settings(con).get("packing_pallet_tare")
    finally:
        con.close()


def test_settings_saves_normalized_tare(admin_client, db_module):
    resp = post_with_csrf(admin_client, "/settings", _settings_form(admin_client, packing_pallet_tare="25,5 кг"),
                          csrf_source_url="/settings")
    assert resp.status_code == 302
    assert _stored_tare(db_module) == "25,5"


@pytest.mark.parametrize("bad", ["abc", "-5", "1000"])
def test_settings_refuses_invalid_tare_and_keeps_the_input(admin_client, db_module, bad):
    form = _settings_form(admin_client, packing_pallet_tare=bad, sender_phone="+359 2 999")
    resp = post_with_csrf(admin_client, "/settings", form, csrf_source_url="/settings")
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    assert "трябва да е число от 0 до 500 кг" in body
    assert 'name="packing_pallet_tare"\n                 value="%s"' % bad in body or \
        re.search(r'name="packing_pallet_tare"\s+value="%s"' % re.escape(bad), body)
    assert 'value="+359 2 999"' in body  # останалото въведено също е запазено
    con = db_module.get_db()
    settings = db_module.get_settings(con)
    con.close()
    assert settings.get("packing_pallet_tare") in (None, "20")
    assert settings.get("sender_phone") != "+359 2 999"  # нищо не е записано


def test_old_unnormalized_tare_still_works(admin_client, db_module):
    con = db_module.get_db()
    db_module.save_settings(con, {"packing_pallet_tare": "20 кг"})
    con.commit()
    con.close()
    assert 'data-pallet-tare="20"' in admin_client.get("/pallet/new").get_data(as_text=True)
    _issue_pallets(admin_client, [_rows(("A1", "d1", "2"))])
    res = _pull(admin_client, "1")
    assert res["row"]["net"] == "80" and res["net_source"] == "tare"


def test_pull_drops_materials_net_greater_than_gross(admin_client, db_module, legacy_packing_pull):
    _add_material(db_module, "A1", "75")           # 2 × 75 = 150 > бруто 100
    _add_material(db_module, "A2", "10")           # 2 × 10 = 20 ≤ 100
    _issue_pallets(admin_client, [_rows(("A1", "d1", "2")), _rows(("A2", "d2", "2"))])
    res = _pull(admin_client, "1")
    assert "net" not in res["row"] and res["net_source"] == ""
    assert "по-голямо от брутото на палетна карта" in res["note"]
    res = _pull(admin_client, "2")
    assert res["row"]["net"] == "20" and res["net_source"] == "materials"


def test_net_from_gross_follows_tare_comma_for_whole_gross():
    import routes_pallet_extra as rpe
    assert rpe._net_from_gross("100", "25,5") == "74,5"
    assert rpe._net_from_gross("100.5", "25,5") == "75"
    assert rpe._net_from_gross("1 250,75", "20") == "1230,75"
    assert rpe._net_from_gross("100", "20") == "80"


# ---------------------------------------------------------------- APPFIX-3

def _pallet(client, order_no, client_name="Müller & Söhne GmbH"):
    data = {"client_name": client_name, "gross": "300", "items_format": "orders",
            "items_json": json.dumps([{"order_no": order_no, "pos": "10", "reference": "R",
                                       "reference_desc": "M", "qty": "1"}])}
    resp = post_with_csrf(client, "/pallet/new", data, csrf_source_url="/pallet/new")
    return int(resp.headers["Location"].rstrip("/").rsplit("/", 1)[-1])


def test_same_po_is_symmetric_with_spaces_in_the_po(admin_client, db_module):
    spaced = _pallet(admin_client, "4500 077 001")
    plain = _pallet(admin_client, "4500077001", client_name="MÜLLER & SÖHNE GMBH")
    _pallet(admin_client, "4500 077 002")                      # друго PO
    _pallet(admin_client, "4500077001", client_name="Друг клиент")
    pulls = []
    for doc_id in (spaced, plain):
        loc = admin_client.get("/pallet/%d/packing-same-po" % doc_id).headers["Location"]
        pulls.append(parse_qs(urlsplit(loc).query)["pull"][0].split(","))
    assert pulls[0] == pulls[1] and len(pulls[0]) == 2


# ---------------------------------------------------------------- APPFIX-4/5/6

@pytest.fixture
def clients(db_module):
    ids = {}
    ids["ist1"] = _add(db_module, "cmr", {"consignee_name": "İstanbul Elektrik Dağıtım A.Ş.", "weight": "100"})
    ids["ist2"] = _add(db_module, "cmr", {"consignee_name": "ISTANBUL ELEKTRIK DAĞITIM A.Ş.", "weight": "50"})
    ids["ist3"] = _add(db_module, "pallet", {"client_name": "İstanbul Elektrik Dağıtım A.Ş."})
    ids["mu1"] = _add(db_module, "cmr", {"consignee_name": "Müller & Söhne GmbH", "weight": "10"})
    ids["mu2"] = _add(db_module, "packing", {"receiver_name": "Müller & Söhne GmbH"})
    ids["mu3"] = _add(db_module, "pallet", {"client_name": "MÜLLER & SÖHNE GMBH"})
    ids["epm1"] = _add(db_module, "cmr", {"consignee_name": "Енерго-Про Мрежи АД", "weight": "1"})
    ids["epm2"] = _add(db_module, "cmr", {"consignee_name": "  ЕНЕРГО-ПРО   МРЕЖИ АД ", "weight": "2"})
    ids["epm3"] = _add(db_module, "pallet", {"client_name": "Енерго-Про Мрежи АД"})
    ids["abb1"] = _add(db_module, "cmr", {"consignee_name": "ABB Ltda Brasil", "weight": "3"})
    ids["abb2"] = _add(db_module, "cmr", {"consignee_name": "ABB Power Grids Norway AS", "weight": "4"})
    ids["inv_epm"] = _add(db_module, "invoice_br", {"consignee_name": "енерго-про мрежи ад",
                                                    "items": [{"qty": "1", "unit_price": "5"}]})
    ids["inv_abb"] = _add(db_module, "invoice_br", {"consignee_name": "ABB Ltda Brasil",
                                                    "items": [{"qty": "1", "unit_price": "5"}]})
    return ids


def _report_clients(client, name):
    body = client.get("/reports?from=2026-07&to=2026-07&group=client&client=" + quote(name)).get_data(as_text=True)
    return [re.sub(r"<[^>]+>", "", c).strip()
            for c in re.findall(r'<td data-label="Клиент">(.*?)</td>', body, re.S)]


@pytest.mark.parametrize("name, expected", [
    ("istanbul elektrik dağıtım a.ş.", ["İstanbul Elektrik Dağıtım A.Ş."]),
    ("ISTANBUL ELEKTRIK DAĞITIM A.Ş.", ["İstanbul Elektrik Dağıtım A.Ş."]),
    ("ABB", ["ABB Ltda Brasil", "ABB Power Grids Norway AS"]),
    ("muller", ["Müller &amp; Söhne GmbH"]),
    ("  енерго-про  мрежи ", ["Енерго-Про Мрежи АД"]),
])
def test_report_client_filter_folds_case_and_matches_part(admin_client, clients, name, expected):
    assert _report_clients(admin_client, name) == expected


def test_report_client_link_shows_every_document_of_the_row(admin_client, clients):
    body = admin_client.get("/reports?from=2026-07&to=2026-07&group=client").get_data(as_text=True)
    link = re.search(r'<a href="(/docs\?[^"]+)"[^>]*>Енерго-Про Мрежи АД</a>', body).group(1).replace("&amp;", "&")
    assert parse_qs(urlsplit(link).query)["client"] == ["Енерго-Про Мрежи АД"]
    docs = admin_client.get(link).get_data(as_text=True)
    assert _ids(docs) == {clients["epm1"], clients["epm2"], clients["epm3"]}
    assert "Клиент: <b>Енерго-Про Мрежи АД</b>" in docs
    # етикетът се маха — адресът е без client, останалото е запазено
    chip = re.search(r'<a class="filter-chip" href="([^"]+)"[^>]*>\s*<span>Клиент:', docs).group(1)
    assert "client=" not in chip and "from=2026-07-01" in chip.replace("&amp;", "&")
    assert '<input type="hidden" name="client" value="Енерго-Про Мрежи АД">' in docs
    inv = admin_client.get("/invoices?client=" + quote("ЕНЕРГО-ПРО МРЕЖИ АД")).get_data(as_text=True)
    assert _ids(inv) == {clients["inv_epm"]} and "Клиент: <b>ЕНЕРГО-ПРО МРЕЖИ АД</b>" in inv


def test_docs_client_filter_is_exact_key_and_combines_with_search(admin_client, clients):
    body = admin_client.get("/docs?client=" + quote("müller & söhne gmbh")).get_data(as_text=True)
    assert _ids(body) == {clients["mu1"], clients["mu2"], clients["mu3"]}
    body = admin_client.get("/docs?client=ABB").get_data(as_text=True)  # не е част от името
    assert _ids(body) == set()
    body = admin_client.get("/docs?type=cmr&client=" + quote("Müller & Söhne GmbH")).get_data(as_text=True)
    assert _ids(body) == {clients["mu1"]}
    # износът в Excel следва филтъра
    from openpyxl import load_workbook
    import io
    resp = admin_client.get("/docs/export.xlsx?client=" + quote("Müller & Söhne GmbH"))
    rows = list(load_workbook(io.BytesIO(resp.data)).active.iter_rows(values_only=True))
    assert len([r for r in rows[1:] if r[1]]) >= 3


def test_group_by_client_uses_the_client_key_and_most_frequent_title(admin_client, clients):
    body = admin_client.get("/docs?group=client").get_data(as_text=True)
    heads = [re.sub(r"<[^>]+>", "", h).strip()
             for h in re.findall(r'class="list-group-row"><th colspan="7">(.*?)</th>', body, re.S)]
    assert heads.count("Müller &amp; Söhne GmbH") == 1 and "MÜLLER &amp; SÖHNE GMBH" not in heads
    assert len([h for h in heads if "stanbul" in h.lower() or "STANBUL" in h]) == 1
    assert len([h for h in heads if "ЕНЕРГО" in h.upper()]) == 1
    # всички документи на групата са под едно заглавие, едно след друго
    order = [int(x) for x in re.findall(r'data-row-id="(\d+)"', body)]
    mu = [order.index(clients[k]) for k in ("mu1", "mu2", "mu3")]
    assert max(mu) - min(mu) == 2


def test_client_match_key(db_module):
    k = db_module.client_match_key
    assert k("  ЕНЕРГО-ПРО   МРЕЖИ АД ") == k("Енерго-Про Мрежи АД")
    assert k("MÜLLER & SÖHNE GMBH") == k("Müller & Söhne GmbH") == "muller & sohne gmbh"
    assert k("İstanbul Elektrik Dağıtım A.Ş.") == k("ISTANBUL ELEKTRIK DAĞITIM A.Ş.")
    assert k("Йордан") != k("Иордан")


# ---------------------------------------------------------------- APPFIX-7/8

def test_system_page_subtitle_has_no_public_address(admin_client):
    body = admin_client.get("/admin/system").get_data(as_text=True)
    assert "Мрежа, архив и възстановяване — общи за цялата инсталация." in body
    assert "публичен адрес —" not in body


def test_code_and_paths_wrap_anywhere():
    css = read_source("static", "style.css")
    assert re.search(r"code, kbd, samp, \.path \{ overflow-wrap: anywhere; \}", css)


# ---------------------------------------------------------------- APPFIX-9 (JS — виж _e2e)

def test_app_js_decimal_separator_handles_grouped_numbers_and_form_context():
    js = read_source("static", "app.js")
    body = js[js.index("function decimalSepOf("):js.index("function withDecimalSep(")]
    assert 'replace(/\\s+/g, "")' in body and "function formDecimalSep(" in body
    assert "decimalSepOf([grossText], decimalSepOf([tareText]))" in js
    assert js.count("formDecimalSep(form") >= 4


# ---------------------------------------------------------------- APPFIX-10

@pytest.fixture
def accented(db_module):
    return {
        "mu": _add(db_module, "cmr", {"consignee_name": "Müller & Söhne GmbH", "consignee_city": "München"}),
        "sp": _add(db_module, "cmr", {"consignee_name": "ABB Ltda", "consignee_city": "Osasco – São Paulo"}),
        "gr": _add(db_module, "packing", {"receiver_name": "Teile", "items": [{"description": "Trennschalter Größe 2"}]}),
        "ep": _add(db_module, "cmr", {"consignee_name": "Енерго-Про Мрежи АД"}),
        "yo": _add(db_module, "cmr", {"consignee_name": "Йордан Йорданов ЕООД"}),
        "io": _add(db_module, "cmr", {"consignee_name": "Иордан Иорданов ЕООД"}),
        "tr": _add(db_module, "cmr", {"consignee_name": "ŞİŞLİ Çağdaş"}),
        "inv": _add(db_module, "invoice_dubai", {"consignee_name": "Søren Ærø São Tomé"}),
    }


def _found(client, query, url="/docs"):
    return _ids(client.get(url, query_string={"q": query}).get_data(as_text=True))


_ACCENT_QUERIES = ["muller", "MULLER", "Müller", "münchen", "munchen", "sao paulo", "SÃO PAULO",
                   "grösse", "grosse", "GRÖSSE", "größe", "енерго про", "про енерго", "енерго мрежи ад",
                   "йордан", "иордан", "ЙОРДАН", "sisli", "şişli", "cagdas", "çağdaş", "sao tome",
                   "søren", "soren", "енерго zzz", "a b", "ab cd", "  муллер  "]


@pytest.mark.parametrize("url", ["/docs", "/invoices"])
def test_accent_insensitive_search_and_same_result_without_index(admin_client, accented, monkeypatch, url):
    import search_index
    new = {q: _found(admin_client, q, url) for q in _ACCENT_QUERIES}
    monkeypatch.setattr(search_index, "_fts_ready", False)
    assert {q: _found(admin_client, q, url) for q in _ACCENT_QUERIES} == new
    if url == "/docs":
        assert new["muller"] == new["MULLER"] == new["Müller"] == {accented["mu"]}
        assert new["münchen"] == new["munchen"] == {accented["mu"]}
        assert new["sao paulo"] == new["SÃO PAULO"] == {accented["sp"]}
        assert new["grösse"] == new["grosse"] == new["GRÖSSE"] == new["größe"] == {accented["gr"]}
        assert new["енерго про"] == new["про енерго"] == new["енерго мрежи ад"] == {accented["ep"]}
        assert new["енерго zzz"] == set()
        # „й“ не е „и“ с диакритика — кирилицата не се сгъва
        assert new["йордан"] == new["ЙОРДАН"] == {accented["yo"]} and new["иордан"] == {accented["io"]}
        assert new["sisli"] == new["şişli"] == new["cagdas"] == new["çağdaş"] == {accented["tr"]}
        assert new["soren"] == set()  # „ø“ не е буква с диакритика (NFKD) — остава
    else:
        assert new["sao tome"] == {accented["inv"]} and new["søren"] == {accented["inv"]}


def test_multi_word_search_uses_the_index_for_all_words(admin_client, accented, monkeypatch):
    import appcore
    import search_index
    sql, params = appcore.document_text_search("енерго про", ("d.number",))
    assert appcore._FTS_SEARCH_SQL in sql and "instr(s2.body, ?)" in sql
    assert params[1:] == ["енерго", "про"]
    assert '"ене"' in params[0] and '"про"' in params[0]
    monkeypatch.setattr(search_index, "_fts_ready", False)
    sql, params = appcore.document_text_search("енерго про", ("d.number",))
    assert sql.count("document_search s WHERE") == 2 and " AND " in sql


def test_search_fold_strips_latin_marks_but_keeps_cyrillic(db_module):
    f = db_module.search_fold
    assert f("Müller São Größe Çağdaş İstanbul") == "muller sao grosse cagdas istanbul"
    assert f("Йордан Ёлка Ѝ") == "йордан ёлка ѝ"
    assert f("й") == "й"   # разложено „й“ → NFC „й“, не „и“
    c = db_module._ci_contains
    long_text = "x" * 300 + "Größe Müller São İstanbul"
    for needle in ("grosse", "muller", "sao", "istanbul", "gros", "GRÖSSE"):
        assert c(long_text, needle), needle
        assert c("Größe Müller São İstanbul", needle), needle
    assert not c("x" * 300 + "Йордан", "иордан") and not c("Йордан", "иордан")


def test_search_index_body_is_folded_and_matches_python_fold(db_module):
    import search_index
    con = db_module.get_db()
    search_index.ensure_schema(con)
    text = "ÀÉÎÕÜ àéîõü Çç Ğğ Şş Ææ Øø Åå Ññ Ăă Șș Țț Łł ẞß İı Йй Ёё Ѝѝ Straße"
    cur = con.execute("INSERT INTO documents (doc_type, number, year, seq, barcode, data)"
                      " VALUES ('cmr', 'N-1', 2026, 1, 'B-1', ?)",
                      (json.dumps({"x": text}, ensure_ascii=False),))
    ascii_id = con.execute("INSERT INTO documents (doc_type, number, year, seq, barcode, data)"
                           " VALUES ('cmr', 'N-2', 2026, 2, 'B-2', ?)",
                           (json.dumps({"x": "Plain ТЕКСТ"}, ensure_ascii=False),)).lastrowid
    escaped = con.execute("INSERT INTO documents (doc_type, number, year, seq, barcode, data)"
                          " VALUES ('cmr', 'N-3', 2026, 3, 'B-3', ?)",
                          (json.dumps({"x": "Müller"}),)).lastrowid  # ü в суровия JSON
    con.commit()
    body = con.execute("SELECT body FROM document_search WHERE id = ?", (cur.lastrowid,)).fetchone()[0]
    folded = body.split(search_index.SEP)[2]
    assert folded == "aeiou aeiou cc gg ss ææ øø aa nn aa ss tt łł ssss ii йй ёё ѝѝ strasse"
    for word in text.split():
        if search_index.is_foldable(word):
            assert search_index.fold_query(word) in folded, word
    assert con.execute("SELECT body FROM document_search WHERE id = ?", (ascii_id,)).fetchone()[0] \
        .endswith("plain текст")
    assert con.execute("SELECT body FROM document_search WHERE id = ?", (escaped,)).fetchone()[0] \
        .endswith("muller")
    con.close()


def test_search_index_is_rebuilt_once_for_an_old_database(db_module):
    """Стара база (тригери/тяло без диакритика, без изгледите) → ensure_schema
    създава изгледите, сменя тригерите и преизгражда тялото; следващото
    стартиране не преизгражда нищо."""
    import search_index
    con = db_module.get_db()
    search_index.ensure_schema(con)
    con.execute("INSERT INTO documents (doc_type, number, year, seq, barcode, data)"
                " VALUES ('cmr', 'O-1', 2026, 1, 'O-1', ?)", (json.dumps({"x": "Müller"}, ensure_ascii=False),))
    for name in list(search_index._FOLD_VIEWS):
        con.execute("DROP VIEW %s" % name)
    con.execute("DROP TRIGGER trg_document_search_update")
    con.execute("CREATE TRIGGER trg_document_search_update AFTER UPDATE OF data ON documents BEGIN SELECT 1; END")
    con.execute("UPDATE document_search SET body = 'o-1\x1fo-1\x1fmüller'")
    con.commit()
    search_index.ensure_schema(con)
    assert con.execute("SELECT body FROM document_search").fetchone()[0] == "o-1\x1fo-1\x1fmuller"
    calls = []
    real = search_index.rebuild
    search_index.rebuild = lambda c: calls.append(1) or real(c)
    try:
        search_index.ensure_schema(con)
    finally:
        search_index.rebuild = real
    assert calls == []
    con.close()


def test_clients_and_materials_search_without_diacritics(admin_client, db_module):
    con = db_module.get_db()
    con.execute("INSERT INTO clients (name, city) VALUES ('Müller & Söhne GmbH', 'München')")
    con.execute("INSERT INTO materials (code, description, net_weight) VALUES ('GCE-1', 'Trennschalter Größe 2', '28')")
    con.commit()
    con.close()
    assert "Müller &amp; Söhne GmbH" in admin_client.get("/clients?q=muller").get_data(as_text=True)
    assert "Trennschalter Größe 2" in admin_client.get("/materials?q=grosse").get_data(as_text=True)


# ---------------------------------------------------------------- APPFIX-11

def test_reversed_dates_are_swapped_in_docs_and_invoices(admin_client, db_module):
    cmr = _add(db_module, "cmr", {"consignee_name": "X"}, created_at="2026-07-15 10:00:00")
    inv = _add(db_module, "invoice_br", {"consignee_name": "X", "items": []}, created_at="2026-07-15 10:00:00")
    _add(db_module, "cmr", {"consignee_name": "X"}, created_at="2026-08-15 10:00:00")
    body = admin_client.get("/docs?from=2026-07-31&to=2026-07-01").get_data(as_text=True)
    assert _ids(body) == {cmr}
    assert 'name="from" value="2026-07-01"' in body and 'name="to" value="2026-07-31"' in body
    body = admin_client.get("/invoices?from=2026-07-31&to=2026-07-01").get_data(as_text=True)
    assert _ids(body) == {inv}


def test_sqlite_views_need_no_app_functions(db_module):
    """Тригерите/изгледите на тялото за търсене ползват само вградени функции —
    връзка без функциите на приложението (друг компютър, скрипт) пак записва."""
    import search_index
    con = db_module.get_db()
    search_index.ensure_schema(con)
    con.close()
    raw = sqlite3.connect(db_module.DB_PATH)
    raw.execute("INSERT INTO documents (doc_type, number, year, seq, barcode, data)"
                " VALUES ('cmr', 'R-1', 2026, 9, 'R-1', ?)", (json.dumps({"x": "Ünal"}, ensure_ascii=False),))
    raw.commit()
    assert raw.execute("SELECT body FROM document_search ORDER BY id DESC").fetchone()[0].endswith("unal")
    raw.close()
