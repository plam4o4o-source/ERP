# -*- coding: utf-8 -*-
"""Регресионни тестове за одита от 26.09.2026 — Excel импортите (поръчки →
палетни карти, редове за фактура, справочник материали), търсенето на
кирилски кодове в справочника и числата от адреса/формата.

Находки: №1 (грешен <dimension> орязва листа), №2 (ред без Order No
изчезва тихо), №3 (кирилски код на материал никога не съвпада), №4
(лъжливо „повече от 5000 реда“ заради празни форматирани редове), №5
(празна разделителна колона печели пред колона „Pallet“), №6 (фактурата
не приема краткия номер на палетна карта), №7 („inf“ в колоната за палет
сваля импорта), №8 (огромни/странни числа в адреса и формата).
"""
import io
import json
import re
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

import invoice_clients_module
import materials
import routes_invoices
import routes_pallet_extra
from conftest import post_with_csrf

_HUGE = "99999999999999999999999"


# ---------------------------------------------------------------- помощни

def _xlsx(rows, sheet_cb=None):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    if sheet_cb:
        sheet_cb(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _patch_dimension(blob, ref):
    """Подменя <dimension ref=…> във всички листове — както го записват
    някои генератори (не самият Excel)."""
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in src.namelist():
            data = src.read(n)
            if n.startswith("xl/worksheets/sheet"):
                data, count = re.subn(rb'<dimension ref="[^"]*"\s*/>',
                                      b'<dimension ref="%s"/>' % ref, data)
                assert count == 1
            z.writestr(n, data)
    return out.getvalue()


def _format_blank_rows(upto):
    """Празни, но форматирани (удебелен шрифт) редове до ред `upto`."""
    def cb(ws):
        for r in range(ws.max_row + 1, upto + 1):
            ws.cell(row=r, column=1).font = Font(bold=True)
    return cb


def _post_file(client, url, blob, src, extra=None):
    data = {"excel_file": (io.BytesIO(blob), "a.xlsx")}
    if extra:
        data.update(extra)
    return post_with_csrf(client, url, data, csrf_source_url=src,
                          content_type="multipart/form-data", follow_redirects=False)


_ORDERS = [["Order No", "Pos", "Reference", "Reference Desc", "Open Qty", None],
           ["4700", "10", "A1", "d1", 5, 1],
           ["4700", "20", "A2", "d2", 6, 2],
           ["4700", "30", "A3", "d3", 7, 2]]

_CATALOG = [["ABB part ID", "Description", "Net weight"],
            ["M1", "d", 1.5], ["M2", "e", 2.5], ["M3", "f", 3.5]]


def _pallet_cards(body):
    """(брой карти, брой редове) от съобщението след bulk импорта."""
    m = re.search(r"Открити са (\d+) палетни карти \((\d+) реда общо\)", body)
    assert m, body[:500]
    return int(m.group(1)), int(m.group(2))


# ---------------------------------------------------------------- №1 грешен <dimension>

def test_wrong_dimension_does_not_truncate_invoice_import(admin_client):
    blob = _patch_dimension(_xlsx(_ORDERS), b"A1:E2")
    res = _post_file(admin_client, "/invoice/import-items", blob, "/invoice-br/new").get_json()
    assert res["ok"], res
    assert res["count"] == 3  # преди поправката: 1


def test_dimension_a1_does_not_reject_valid_invoice_file(admin_client):
    blob = _patch_dimension(_xlsx(_ORDERS), b"A1")
    res = _post_file(admin_client, "/invoice/import-items", blob, "/invoice-br/new").get_json()
    assert res["ok"], res  # преди: „не съдържа разпознаваеми колони“
    assert res["count"] == 3


def test_wrong_dimension_does_not_truncate_pallet_import_and_keeps_grouping(admin_client):
    """Освен пълния брой редове проверява и, че безименната групираща
    колона оцелява след reset_dimensions() — клетката ѝ в заглавния ред
    изобщо не е записана във файла, тоест заглавието е по-късо от
    редовете (без допълването в _parse_order_export всичко отиваше по
    колона „Open Qty“)."""
    blob = _patch_dimension(_xlsx(_ORDERS), b"A1:E2")
    resp = _post_file(admin_client, "/pallet/bulk-import", blob, "/pallet/new")
    assert resp.status_code == 200, resp.headers.get("Location")
    assert _pallet_cards(resp.get_data(as_text=True)) == (2, 3)


def test_wrong_dimension_does_not_truncate_materials_import(db_module):
    entries = materials.parse_catalog_xlsx(_patch_dimension(_xlsx(_CATALOG), b"A1:B2"))
    assert entries is not None
    assert [e[0] for e in entries] == ["M1", "M2", "M3"]
    assert [e[2] for e in entries] == ["1.5", "2.5", "3.5"]  # колона C извън ref
    assert materials.parse_catalog_xlsx(_patch_dimension(_xlsx(_CATALOG), b"A1")) is not None


# ---------------------------------------------------------------- №2 ред без Order No

def test_rows_without_order_no_are_kept_with_a_warning():
    wb = Workbook()
    ws = wb.active
    for r in (["Order No", "Pos", "Reference", "Reference Desc", "Open Qty", ""],
              ["4700", "10", "A1", "d1", 5, 1],
              [None, "20", "A2", "d2", 6, 1],
              [None, "30", "A3", "d3", 7, 2],
              [None, None, None, None, None, 2]):  # само номер на палет — не е артикул
        ws.append(r)
    groups, warnings = routes_pallet_extra._parse_order_export(ws)
    refs = sorted(it["reference"] for items in groups.values() for it in items)
    assert refs == ["A1", "A2", "A3"]  # преди: само A1
    # Номерът НЕ се досеща/попълва надолу.
    assert [it["order_no"] for it in groups[2]] == [""]
    assert any("2 ред(а) без номер на поръчка (Order No)" in w for w in warnings), warnings


def test_rows_without_order_no_reach_the_bulk_review(admin_client):
    blob = _xlsx([["Order No", "Pos", "Reference", "Reference Desc", "Open Qty", None],
                  ["4700", "10", "A1", "d1", 5, 1],
                  [None, "20", "A2", "d2", 6, 1]])
    resp = _post_file(admin_client, "/pallet/bulk-import", blob, "/pallet/new")
    assert resp.status_code == 200
    body = resp.get_data(as_text=True)
    assert _pallet_cards(body) == (1, 2)
    assert "без номер на поръчка (Order No)" in body


# ---------------------------------------------------------------- №3 кирилски кодове

def test_lookup_many_matches_cyrillic_codes_in_any_case(db_module):
    con = db_module.get_db()
    materials.replace_catalog(con, [("Кабел-1", "Кабел", "1.5"), ("ABC-1", "x", "2")])
    for code in ("Кабел-1", "кабел-1", "КАБЕЛ-1", "кабел-1-RAS"):
        found = materials.lookup_many(con, [code])
        assert code in found, code
        assert found[code]["net_weight"] == "1.5"
    # ASCII пътят е непроменен.
    assert materials.lookup_many(con, ["abc-1"])["abc-1"]["code"] == "ABC-1"
    assert materials.lookup(con, "кабел-1")["code"] == "Кабел-1"
    assert materials.lookup(con, "КАБЕЛ-1-RAS")["code"] == "Кабел-1"


def test_lookup_many_cyrillic_is_chunked(db_module):
    con = db_module.get_db()
    materials.replace_catalog(con, [("Код-%d" % i, "", str(i)) for i in range(1200)])
    found = materials.lookup_many(con, ["код-%d" % i for i in range(1200)])
    assert len(found) == 1200
    assert found["код-1199"]["net_weight"] == "1199"


def test_materials_lookup_route_and_invoice_import_with_cyrillic_code(admin_client, db_module):
    materials.replace_catalog(db_module.get_db(), [("Кабел-1", "Кабел", "1.5")])
    res = admin_client.get("/materials/lookup?code=кабел-1").get_json()
    assert res["ok"] and res["code"] == "Кабел-1"
    blob = _xlsx([["Order No", "Pos", "Reference", "Reference Desc", "Open Qty"],
                  ["4700", "10", "Кабел-1", "", 5]])
    res = _post_file(admin_client, "/invoice/import-items", blob, "/invoice-br/new").get_json()
    assert res["matched"] == 1, res  # преди: 0
    assert res["rows"][0]["net_weight"] == "1.5"


# ---------------------------------------------------------------- №4 лъжливо орязване

_TRUNC = "Файлът съдържа повече от"


@pytest.fixture
def small_cap(monkeypatch):
    """Таван от 5 реда вместо 5000 — същата логика, без 6000-редов файл."""
    for mod in (routes_invoices, routes_pallet_extra, materials):
        monkeypatch.setattr(mod, "_MAX_IMPORT_DATA_ROWS", 5)
    return 5


def test_formatted_empty_rows_do_not_trigger_truncation_warning_invoice(admin_client, small_cap):
    blob = _xlsx(_ORDERS, _format_blank_rows(200))
    res = _post_file(admin_client, "/invoice/import-items", blob, "/invoice-br/new").get_json()
    assert res["ok"], res
    assert not any(_TRUNC in w for w in res.get("warnings", [])), res.get("warnings")


def test_formatted_empty_rows_do_not_trigger_truncation_warning_pallet(admin_client, small_cap):
    blob = _xlsx(_ORDERS, _format_blank_rows(200))
    resp = _post_file(admin_client, "/pallet/bulk-import", blob, "/pallet/new")
    assert resp.status_code == 200
    assert _TRUNC not in resp.get_data(as_text=True)


def test_formatted_empty_rows_do_not_trigger_truncation_materials(db_module, small_cap):
    stats = {}
    entries = materials.parse_catalog_xlsx(_xlsx(_CATALOG, _format_blank_rows(200)), stats)
    assert len(entries) == 3
    assert stats["truncated"] is False
    # И точно 5 реда данни + празни форматирани редове след тях.
    rows = [_CATALOG[0]] + [["M%d" % i, "d", 1] for i in range(5)]
    stats = {}
    assert len(materials.parse_catalog_xlsx(_xlsx(rows, _format_blank_rows(200)), stats)) == 5
    assert stats["truncated"] is False


def test_real_data_after_empty_rows_still_triggers_truncation(admin_client, db_module, small_cap):
    """Обратната страна — поправката не бива да скрие истинско орязване,
    дори данните след тавана да са отделени с празни форматирани редове."""
    def cb(ws):
        _format_blank_rows(100)(ws)
        ws.cell(row=101, column=1, value="4700")
        ws.cell(row=101, column=3, value="LATE")
        ws.cell(row=101, column=5, value=1)

    rows = _ORDERS + [["4700", str(i), "B%d" % i, "", 1, 1] for i in range(4)]
    res = _post_file(admin_client, "/invoice/import-items", _xlsx(rows, cb),
                     "/invoice-br/new").get_json()
    assert any(_TRUNC in w for w in res.get("warnings", []))
    resp = _post_file(admin_client, "/pallet/bulk-import", _xlsx(rows, cb), "/pallet/new")
    assert _TRUNC in resp.get_data(as_text=True)
    cat = [_CATALOG[0]] + [["M%d" % i, "d", 1] for i in range(5)]
    stats = {}
    materials.parse_catalog_xlsx(_xlsx(cat, cb), stats)
    assert stats["truncated"] is True


def test_truncation_scan_is_bounded_and_errs_on_the_side_of_warning(monkeypatch):
    monkeypatch.setattr(materials, "_TRUNCATION_SCAN_ROWS", 3)
    assert materials.more_data_follows(iter([(None,), ("",), (" ",)])) is False
    assert materials.more_data_follows(iter([(None,), (None, "x")])) is True
    # След прегледаните 3 празни реда има още — по-добре предупреждение.
    assert materials.more_data_follows(iter([(None,)] * 4)) is True


# ---------------------------------------------------------------- №5 групираща колона

def test_named_pallet_column_wins_over_empty_spacer_column():
    wb = Workbook()
    ws = wb.active
    for r in (["Order No", None, "Pos", "Reference", "Reference Desc", "Open Qty", "Pallet"],
              ["4700", None, "10", "A1", "d", 5, 1],
              ["4700", None, "20", "A2", "d", 5, 2]):
        ws.append(r)
    groups, warnings = routes_pallet_extra._parse_order_export(ws)
    assert sorted(groups) == [1, 2]  # преди: всичко в карта №1
    assert not any("без разчетен номер на палет" in w for w in warnings), warnings


def test_named_pallet_column_wins_even_when_it_is_not_the_last_column():
    wb = Workbook()
    ws = wb.active
    for r in (["Order No", None, "Pos", "Reference", "Reference Desc", "Open Qty",
               "Pallet No", "Stock"],
              ["4700", None, "10", "A1", "d", 5, 1, 40],
              ["4700", None, "20", "A2", "d", 5, 2, 50]):
        ws.append(r)
    groups, _w = routes_pallet_extra._parse_order_export(ws)
    assert sorted(groups) == [1, 2]  # не 40/50 от „Stock“


def test_unrecognized_last_named_column_wins_over_empty_spacer_column():
    wb = Workbook()
    ws = wb.active
    for r in (["Order No", None, "Pos", "Reference", "Reference Desc", "Open Qty", "PAL"],
              ["4700", None, "10", "A1", "d", 5, 1],
              ["4700", None, "20", "A2", "d", 5, 2]):
        ws.append(r)
    groups, _w = routes_pallet_extra._parse_order_export(ws)
    assert sorted(groups) == [1, 2]


# ---------------------------------------------------------------- №6 кратък номер във фактурата

def test_invoice_pull_pallet_accepts_the_short_number(admin_client, db_module):
    items = [{"order_no": "4700", "pos": "10", "reference": "A1",
              "reference_desc": "d", "qty": "3"}]
    resp = post_with_csrf(admin_client, "/pallet/new",
                          {"client_name": "K", "items_format": "orders",
                           "items_json": json.dumps(items)},
                          csrf_source_url="/pallet/new", follow_redirects=False)
    assert resp.status_code == 302
    number = db_module.get_db().execute(
        "SELECT number FROM documents WHERE doc_type = 'pallet' ORDER BY id DESC").fetchone()[0]
    short = str(int(number.split("/")[0]))
    res = post_with_csrf(admin_client, "/invoice/pull-pallet", {"code": short},
                         csrf_source_url="/invoice-br/new").get_json()
    assert res.get("ok"), res  # преди: „Няма документ с номер/баркод“
    assert res["rows"][0]["material_code"] == "A1"


# ---------------------------------------------------------------- №7 inf/nan в колоната за палет

@pytest.mark.parametrize("raw", ["inf", "-inf", "nan", "1e999"])
def test_group_number_inf_nan_falls_back_to_card_one(raw):
    assert routes_pallet_extra._parse_group_numbers(raw) == ([1], True)


def test_bulk_import_with_inf_group_does_not_crash(admin_client):
    blob = _xlsx([["Order No", "Pos", "Reference", "Reference Desc", "Open Qty", None],
                  ["4700", "10", "ABC", "d", 5, "inf"]])
    resp = _post_file(admin_client, "/pallet/bulk-import", blob, "/pallet/new")
    assert resp.status_code == 200, resp.headers.get("Location")  # преди: 302
    assert "без разчетен номер на палет" in resp.get_data(as_text=True)


# ---------------------------------------------------------------- №8 огромни числа

@pytest.mark.parametrize("url", [
    "/invoices?page=" + _HUGE, "/invoices?page=-3", "/invoices?page=0",
    "/invoices/clients?page=" + _HUGE,
])
def test_huge_or_negative_page_does_not_crash(admin_client, url):
    assert admin_client.get(url).status_code == 200


def test_page_arg_is_clamped(flask_app):
    with flask_app.test_request_context("/?page=" + _HUGE):
        assert routes_invoices._page_arg() == routes_invoices._MAX_PAGE
    with flask_app.test_request_context("/?page=-7"):
        assert routes_invoices._page_arg() == 1


def test_parse_id_list_drops_out_of_range_and_non_decimal():
    assert routes_pallet_extra._parse_id_list(
        "1, 2,%s,²,0,-3,x,%d,%d" % (_HUGE, 2 ** 63 - 1, 2 ** 63)) == [1, 2, 2 ** 63 - 1]


def test_bulk_print_with_one_huge_id_still_prints_the_valid_ones(admin_client, db_module):
    resp = post_with_csrf(admin_client, "/pallet/new", {"client_name": "K"},
                          csrf_source_url="/pallet/new", follow_redirects=False)
    assert resp.status_code == 302
    doc_id = db_module.get_db().execute(
        "SELECT id FROM documents WHERE doc_type = 'pallet' ORDER BY id DESC").fetchone()[0]
    resp = admin_client.get("/pallet/bulk-print?ids=%d,%s" % (doc_id, _HUGE))
    assert resp.status_code == 200  # преди: 302 (OverflowError)
    for ids in (_HUGE, "²"):
        assert admin_client.get("/pallet/bulk-result?ids=" + ids).status_code == 200


@pytest.mark.parametrize("url,src", [("/packing/pull-pallet", "/packing/new"),
                                     ("/invoice/pull-pallet", "/invoice-br/new")])
def test_pull_pallet_with_superscript_digit_does_not_crash(admin_client, url, src):
    resp = post_with_csrf(admin_client, url, {"code": "²"}, csrf_source_url=src)
    assert resp.status_code == 200
    assert resp.get_json()["ok"] is False


def test_invoice_client_get_with_out_of_range_id(db_module):
    con = db_module.get_db()
    assert invoice_clients_module.get(con, 2 ** 70) is None
    assert invoice_clients_module.get(con, 0) is None
