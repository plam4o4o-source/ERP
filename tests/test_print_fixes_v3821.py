# -*- coding: utf-8 -*-
"""Одит (21.10.2026, PRINT-1…14): печатни бланки и PDF износ (v3.82.1).
Измерванията в истински печатен движок са в test_print_fixes_v3821_e2e.py."""
import io
import json
import re

import pytest

from conftest import get_csrf_token, read_source
from test_print_layout import _doc

pypdf = pytest.importorskip("pypdf")

PRINT_TEMPLATES = ("waybill_print.html", "pallet_bulk_preview.html", "dualuse_print.html",
                   "pallet_bulk_print.html", "packing_print.html", "cmr_print.html",
                   "pallet_print.html", "invoice_no_print.html", "export_it_print.html",
                   "invoice_br_print.html", "shipment_print.html", "invoice_dubai_print.html")


def _css():
    return read_source("static", "style.css")


def _html(client, url):
    resp = client.get(url)
    assert resp.status_code == 200, url
    return resp.get_data(as_text=True)


def _pdf_reader(client, doc_id):
    resp = client.get("/doc/%d/export.pdf" % doc_id)
    assert resp.status_code == 200 and resp.mimetype == "application/pdf"
    return pypdf.PdfReader(io.BytesIO(resp.data))


# ---------------------------------------------------------------- PRINT-1 / PRINT-4 / PRINT-5
def test_waybill_goods_cells_wrap_long_codes():
    css = _css()
    assert ".twb table.twb-goods th, .twb table.twb-goods td { overflow-wrap: anywhere; }" in css


def test_row_number_column_never_wraps():
    css = _css()
    m = re.search(r"\.pkl table\.goods \.rn, \.inv table\.goods \.rn, \.twb table\.twb-goods \.rn \{([^}]*)\}", css)
    assert m and "white-space: nowrap" in m.group(1) and "overflow-wrap: normal" in m.group(1)
    for name in ("packing_print.html", "invoice_br_print.html", "invoice_no_print.html",
                 "invoice_dubai_print.html", "waybill_print.html"):
        tpl = read_source("templates", name)
        assert '<td class="c rn">{{ loop.index }}</td>' in tpl, name
        assert re.search(r'<th class="rn" style="width:\d+px">№</th>', tpl), name


def test_packing_meta_line_wraps():
    css = _css()
    assert ".pkl-meta { flex-wrap: wrap; row-gap: 4px; }" in css
    assert ".pkl-meta > div { min-width: 0; overflow-wrap: anywhere; }" in css


# ---------------------------------------------------------------- PRINT-2 / PRINT-3: палетна карта
def test_pallet_label_has_extra_font_steps_and_js_fit(admin_client, db_module):
    js = read_source("static", "app.js")
    assert "function fitPalletLabels()" in js and "initPalletLabelFit();" in js
    fit = js[js.index("function fitPalletLabels"):js.index("function initPalletLabelFit")]
    assert ".plt-client-box, .plt-stats" in fit and "style.zoom" in fit
    assert "plt-head" not in fit  # баркодът не се смалява
    long = "Акционерно дружество " * 12
    doc_id = _doc(db_module, "pallet", {"client_name": long, "client_address": long, "net": "1",
                                        "items_format": "orders", "items": []})
    html = _html(admin_client, "/doc/%d?format=label" % doc_id)
    assert "plt-client-xxl" in html
    css = _css()
    assert ".label-format .plt-client-box.plt-client-xxl .val {" in css
    assert ".label-format .plt-head-barcode { flex: none; max-width: none; }" in css


@pytest.mark.parametrize("url", ["/doc/%d", "/pallet/bulk-print?ids=%d"])
def test_pallet_card_cmr_ref_on_its_own_row(admin_client, db_module, url):
    doc_id = _doc(db_module, "pallet", {"client_name": "К", "ref_cmr": "0001/2026, 0002/2026",
                                        "items_format": "orders", "items": []})
    html = _html(admin_client, url % doc_id)
    assert '<div class="plt-date plt-ref">ЧМР / CMR №: <b>0001/2026, 0002/2026</b></div>' in html
    assert ".plt-head-barcode { flex: 0 0 107mm; max-width: 58%; }" in _css()


# ---------------------------------------------------------------- PRINT-6: lang на бланките
def test_every_print_page_is_marked_bilingual_bg():
    for name in PRINT_TEMPLATES:
        tpl = read_source("templates", name)
        tags = re.findall(r'<div class="print-page[^>]*>', tpl)
        assert tags, name
        assert all('lang="bg"' in t for t in tags), (name, tags)
    js = read_source("static", "app.js")
    fit = js[js.index("function fitWaybillPages"):js.index("function initWaybillPrintFit")]
    assert 'page.setAttribute("lang", first.getAttribute("lang") || "bg");' in fit


# ---------------------------------------------------------------- PRINT-7: заглавен лист
def test_shipment_cover_total_is_last_body_row(admin_client, db_module):
    a = _doc(db_module, "cmr", {"consignee_name": "К", "packages": "2"})
    b = _doc(db_module, "packing", {"receiver_name": "К", "items": [{"description": "x", "qty": "1"}]})
    html = _html(admin_client, "/shipment/print?ids=%d:1,%d:1&cover=1&from=%d" % (a, b, a))
    cover = html[html.index('class="print-page shipment-cover"'):html.index('<section class="shipment-doc"')]
    assert 'lang="bg"' in cover[:60]
    assert "<tfoot>" not in cover
    body = cover[cover.index("<tbody>"):cover.index("</tbody>")]
    assert body.rstrip().endswith('<td class="n">2</td></tr>') and "Общо / Total" in body


# ---------------------------------------------------------------- PRINT-8 / PRINT-11: PDF износ
def _fields(n=17):
    return [("Поле %d" % i, "Стойност %d" % i) for i in range(n)]


def _inv_cols():
    return [("hs_code", "HS code"), ("po_no", "P.O NO"), ("material_code", "Код"), ("qty", "Количество"),
            ("unit_price", "Цена")]


def test_pdf_fields_are_laid_out_in_a_grid():
    import pdf_export
    rows = pdf_export.pdf_field_grid(
        [("Дата", "20.10.2026"), ("Празно", ""), ("Държава", "Bulgaria"), ("Вид", "Truck"),
         ("Адрес", "Organize Sanayi Bölgesi, 4. Kısım, Atatürk Bulvarı No: 123/A-B"),
         ("Тел.", "+359"), ("Бележка", "Organize Sanayi Bölgesi, 4. Kısım, Atatürk Bulvarı No: 123/A-B, Kat 7")])
    assert [[span for _l, _v, span, _p in r] for r in rows] == [[1, 1, 1], [2, 1], [3]]
    assert rows[0][0][:2] == (["Дата"], ["20.10.2026"])
    assert all(sum(c[2] for c in r) == pdf_export._PDF_FIELD_GRID_COLS for r in rows)


def test_short_invoice_pdf_fits_one_page(flask_app):
    import pdf_export
    items = [dict(hs_code="7326909890", po_no="PO-77881", material_code="MAT-%05d" % i, qty=str(i),
                  unit_price="2.55") for i in range(1, 6)]
    with flask_app.test_request_context():
        pdf = pdf_export.generate_document_pdf("Фактура за Бразилия", "BR-1", "", _fields(), items,
                                               _inv_cols(), totals_row=["TOTAL", "", "", "15", ""])
    assert len(pypdf.PdfReader(io.BytesIO(pdf)).pages) == 1


def test_pdf_total_row_never_stands_alone_on_a_page(flask_app):
    """За всеки брой редове, при който таблицата прелива: последният лист носи
    и поне един ред с данни, не само заглавния ред + TOTAL."""
    import pdf_export
    cols = [("order_no", "Поръчка"), ("reference", "Референция"), ("qty", "Количество")]
    checked = 0
    with flask_app.test_request_context():
        for n in range(30, 62):
            items = [dict(order_no="PO-1", reference="REF-%05d" % i, qty="1") for i in range(n)]
            pdf = pdf_export.generate_document_pdf("Палетна карта", "1", "", _fields(3), items, cols,
                                                   totals_row=["ОБЩО / TOTAL", "", str(n)])
            pages = pypdf.PdfReader(io.BytesIO(pdf)).pages
            if len(pages) < 2:
                continue
            checked += 1
            last = pages[-1].extract_text()
            assert "TOTAL" in last and "REF-%05d" % (n - 1) in last, n
    assert checked >= 5
    assert pdf_export._keep_tail_rows == [0]


def test_invoice_pdf_has_no_barcode_other_types_keep_it(admin_client, db_module):
    inv = _doc(db_module, "invoice_no", {"invoice_number": "NO-1", "items": [
        {"material_code": "M1", "qty": "1", "unit_price": "10"}]})
    cmr = _doc(db_module, "cmr", {"consignee_name": "К"})

    def images(doc_id):
        reader = _pdf_reader(admin_client, doc_id)
        return sum(len(p.images) for p in reader.pages)
    assert images(inv) == 0
    assert images(cmr) == 1


# ---------------------------------------------------------------- PRINT-9: празна дата
def test_packing_without_date_prints_dash(admin_client, db_module):
    doc_id = _doc(db_module, "packing", {"receiver_name": "К", "items": [{"description": "x", "qty": "1"}]})
    assert "Дата / Date: <b>—</b>" in _html(admin_client, "/doc/%d" % doc_id)


# ---------------------------------------------------------------- PRINT-10: Excel числа
@pytest.mark.parametrize("doc_type,key,label,value", [
    ("cmr", "packages", "Брой колети", 12),
    ("waybill", "mileage", "Пробег, км", 220),
    ("waybill", "escort_count", "Брой придружители", 2),
])
def test_excel_numeric_header_fields(admin_client, db_module, doc_type, key, label, value):
    from openpyxl import load_workbook
    doc_id = _doc(db_module, doc_type, {"consignee_name": "К", key: str(value),
                                        "items": [{"description": "x", "qty": "1"}]})
    resp = admin_client.get("/doc/%d/export.xlsx" % doc_id)
    ws = load_workbook(io.BytesIO(resp.data)).active
    cells = [row[1].value for row in ws.iter_rows() if row[0].value == label]
    assert cells == [value], cells


# ---------------------------------------------------------------- PRINT-12: превод
def test_pallet_bulk_print_header_is_translatable(admin_client, db_module):
    ids = [_doc(db_module, "pallet", {"client_name": "К%d" % i, "items_format": "orders", "items": []})
           for i in range(3)]
    html = _html(admin_client, "/pallet/bulk-print?ids=" + ",".join(map(str, ids)))
    assert "Издадени палетни карти, готови за печат: 3" in html
    assert "издадени палетни карти — готови за печат" not in html
    tpl = read_source("templates", "pallet_bulk_print.html")
    assert "_('Издадени палетни карти, готови за печат: %(count)d', count=docs|length)" in tpl
    assert "_('Печат на палетни карти: %(count)d', count=docs|length)" in tpl


# ---------------------------------------------------------------- PRINT-13 / PRINT-14: CSS
def test_shipment_toolbar_wraps_on_phone():
    assert ("main.container-print > .doc-toolbar.shipment-toolbar { flex-wrap: wrap; overflow-x: visible; }"
            in _css())


def test_page_numbers_in_page_margin_except_label_and_cmr():
    css = _css()
    assert re.search(r'@page \{ @bottom-right \{ content: "стр\. " counter\(page\) " / " counter\(pages\);', css)
    assert "@page label { @bottom-right { content: none; } }" in css
    assert "@page cmr { @bottom-right { content: none; } }" in css
    assert ".print-page.cmr-page { page: cmr; }" in css
