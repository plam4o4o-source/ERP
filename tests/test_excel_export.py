# -*- coding: utf-8 -*-
"""Excel износ на документ (/doc/<id>/export.xlsx) — одит 01.10.2026, F3.

Износът питаше `ws.max_row`/`ws[ws.max_row]` след всеки добавен ред, а
openpyxl пресмята това с обхождане на ВСИЧКИ клетки — квадратично време при
стотици редове. Освен това `header_row = ws.max_row + 1` сочеше ПРАЗНИЯ ред
пред таблицата (append([]) не създава клетки), затова удебелен излизаше
празният ред, а не заглавният ред на колоните."""
import io
import json

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from conftest import post_with_csrf


def _issue(client, url, fields):
    resp = post_with_csrf(client, url, fields, csrf_source_url=url, follow_redirects=False)
    assert resp.status_code == 302, resp.status_code
    return int(resp.headers["Location"].rstrip("/").split("/")[-1])


def _packing(client, rows=2):
    items = [{"packing": "кашон", "description": "стока %d" % i, "qty": "2",
              "length": "1", "width": "2", "height": "3", "volume": "0.006",
              "net": "4", "gross": "5"} for i in range(rows)]
    return _issue(client, "/packing/new", {
        "receiver_name": "Получател", "total_net": "8",
        "items_json": json.dumps(items, ensure_ascii=False)})


def _sheet(client, doc_id):
    resp = client.get("/doc/%d/export.xlsx" % doc_id)
    assert resp.status_code == 200
    return load_workbook(io.BytesIO(resp.data)).active


def _row_of(ws, first_value):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == first_value:
            return row[0].row
    raise AssertionError("няма ред, започващ с %r" % first_value)


def test_item_header_row_is_bold_and_the_blank_row_before_it_is_not(admin_client):
    ws = _sheet(admin_client, _packing(admin_client))
    header = _row_of(ws, "Вид опаковка")
    assert all(ws.cell(row=header, column=c).font.bold for c in range(1, 10)), (
        "заглавният ред на таблицата с редове не е удебелен")
    blank = header - 1
    assert all(ws.cell(row=blank, column=c).value is None for c in range(1, 10))
    assert not any(ws.cell(row=blank, column=c).font.bold for c in range(1, 10)), (
        "удебелен е празният ред пред таблицата вместо заглавния")


def test_field_and_item_number_formats_land_on_their_own_rows(admin_client):
    ws = _sheet(admin_client, _packing(admin_client))
    net_row = _row_of(ws, "Общо нето, кг")
    assert ws.cell(row=net_row, column=1).font.bold
    assert ws.cell(row=net_row, column=2).value == 8
    assert ws.cell(row=net_row, column=2).number_format == "0.######"
    first_item = _row_of(ws, "Вид опаковка") + 1
    assert ws.cell(row=first_item, column=3).value == 2
    assert ws.cell(row=first_item, column=3).number_format == "0.######"


def test_invoice_totals_row_is_bold_and_money_formatted(admin_client):
    items = [{"material_code": "A", "qty": "2", "unit_price": "1.5", "net_weight": "0.5"},
             {"material_code": "B", "qty": "1", "unit_price": "=1+1"}]
    doc_id = _issue(admin_client, "/invoice-br/new", {
        "consignee_name": "ABB", "invoice_number": "XL-1",
        "items_json": json.dumps(items)})
    ws = _sheet(admin_client, doc_id)
    first_item = _row_of(ws, "HS code") + 1
    assert ws.cell(row=first_item, column=7).number_format == "0.00###"
    # Стойност като формула остава текст (quotePrefix) на СВОЯ ред.
    formula_cell = ws.cell(row=first_item + 1, column=7)
    assert formula_cell.value == "=1+1" and formula_cell.data_type == "s"
    assert formula_cell.quotePrefix
    total = _row_of(ws, "TOTAL")
    assert all(ws.cell(row=total, column=c).font.bold for c in range(1, 10))


def test_export_does_not_rescan_the_whole_sheet_for_every_row(admin_client, monkeypatch):
    """F3: измерено 306→63 ms при 475 реда. Тук броим обхожданията на листа —
    те не бива да растат с броя на редовете."""
    doc_id = _packing(admin_client, rows=60)
    calls = []
    original = Worksheet.max_row

    def counting(self):
        calls.append(1)
        return original.fget(self)

    monkeypatch.setattr(Worksheet, "max_row", property(counting))
    ws = _sheet(admin_client, doc_id)
    monkeypatch.setattr(Worksheet, "max_row", original)
    assert ws.max_row > 60
    assert len(calls) < 10, "ws.max_row е извикан %d пъти" % len(calls)
