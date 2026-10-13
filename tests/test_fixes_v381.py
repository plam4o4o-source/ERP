# -*- coding: utf-8 -*-
"""Одит (13.10.2026, група FORMS, v3.81.0): спасяване на формите при изтекла
сесия (партида палетни карти, клиент, запис за фактури, материал, фирмени
данни), `next` с параметрите на адреса, задължителните полета на фактурите
и при редакция, текстове/множествено число, „Копирай като нов“ без
препратките към пратката, диапазон от палетни карти, дата в „Компютри и
версии“, „Общо“ в износа на документите и Excel износ на фактурите."""
import html
import io
import json
import re
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook

import appcore
from conftest import get_csrf_token, issue_cmr, post_with_csrf, read_source, with_required

ORIGIN = {"Origin": "http://localhost"}
ADMIN = ("test_admin", "test-password-123")


def _login(client, url, username=ADMIN[0], password=ADMIN[1]):
    token = get_csrf_token(client, url)
    return client.post(url, data={"csrf_token": token, "username": username,
                                  "password": password})


def _doc_count(db_module, doc_type):
    con = db_module.get_db()
    try:
        return con.execute("SELECT COUNT(*) FROM documents WHERE doc_type = ?",
                           (doc_type,)).fetchone()[0]
    finally:
        con.close()


def _doc_data(db_module, doc_id):
    con = db_module.get_db()
    try:
        return json.loads(con.execute("SELECT data FROM documents WHERE id = ?",
                                      (doc_id,)).fetchone()[0])
    finally:
        con.close()


def _issue(client, url, fields):
    resp = post_with_csrf(client, url, with_required(url, fields), csrf_source_url=url)
    assert resp.status_code == 302 and "/doc/" in resp.headers["Location"], resp.headers
    return int(resp.headers["Location"].rstrip("/").rsplit("/", 1)[-1])


def _expired_post(flask_app, admin_client, form_url, post_url, data, **kw):
    """POST с токена на влезлия админ, но от „изтекла“ сесия (нова бисквитка)."""
    token = get_csrf_token(admin_client, form_url)
    fresh = flask_app.test_client()
    resp = fresh.post(post_url, data=dict(data, csrf_token=token), headers=ORIGIN, **kw)
    return fresh, resp


# ============================================================ FORMS-4


def _bulk_form():
    items = [{"order_no": "4700", "pos": "10", "reference": "A1", "qty": "5"}]
    return {"groups": "1,2", "client_name": "Клиент Партида ЕООД", "ref_cmr": "CMR-7781",
            "items_json_1": json.dumps(items), "items_format_1": "orders",
            "items_json_2": json.dumps([dict(items[0], pos="20")]), "items_format_2": "orders",
            "gross_1": "351,5", "height_1": "145", "pallet_type_1": "120×80",
            "gross_2": "222", "height_2": "160", "pallet_type_2": "120×80",
            "issue_token": "tok-forms-4-abcdef"}


def test_bulk_pallet_issue_with_an_expired_session_is_restored_after_login(
        flask_app, db_module, admin_client):
    fresh, resp = _expired_post(flask_app, admin_client, "/pallet/new", "/pallet/bulk-issue",
                                _bulk_form())
    assert resp.status_code == 302, "без поправката: 400 и всичко въведено се губи"
    location = resp.headers["Location"]
    assert location.startswith("/login?") and "rescue=" in location
    assert _doc_count(db_module, "pallet") == 0
    resp = _login(fresh, location)
    assert resp.headers["Location"].startswith("/pallet/bulk-review/restore/")
    body = fresh.get(resp.headers["Location"]).get_data(as_text=True)
    for text in ("351,5", "145", "222", "CMR-7781", "Клиент Партида ЕООД"):
        assert text in body, text
    assert "Попълнената форма е възстановена" in body
    # FORMS-12: едно изречение с числото вътре (не три парчета около него).
    assert "„1 от 2“, „2 от 2“" in html.unescape(body)


def test_bulk_pallet_rescue_is_not_handed_to_another_user(flask_app, db_module, admin_client):
    from werkzeug.security import generate_password_hash
    con = db_module.get_db()
    con.execute("INSERT INTO users (username, password_hash, full_name, role, active,"
                " must_change_password) VALUES ('drug', ?, 'Друг', 'employee', 1, 0)",
                (generate_password_hash("other-pass-987"),))
    con.commit()
    con.close()
    _fresh, resp = _expired_post(flask_app, admin_client, "/pallet/new", "/pallet/bulk-issue",
                                 _bulk_form())
    other = flask_app.test_client()
    resp = _login(other, resp.headers["Location"], "drug", "other-pass-987")
    assert "bulk-review/restore" not in resp.headers["Location"]
    assert "CMR-7781" not in other.get(resp.headers["Location"],
                                       follow_redirects=True).get_data(as_text=True)


# ============================================================ FORMS-5


@pytest.mark.parametrize("path", ["/pallet/bulk-result?ids=1,2", "/shipment/print?ids=3,4",
                                  "/docs?type=cmr&q=%D0%90%D0%BB%D1%84%D0%B0"])
def test_login_redirect_keeps_the_query_string(client, path):
    resp = client.get(path)
    assert resp.status_code == 302
    nxt = parse_qs(urlsplit(resp.headers["Location"]).query)["next"][0]
    assert nxt == urlsplit(path).path + "?" + urlsplit(path).query, nxt


def test_login_follows_next_with_query_and_still_refuses_foreign_hosts(flask_app, admin_client):
    c = flask_app.test_client()
    location = c.get("/pallet/bulk-result?ids=1,2").headers["Location"]
    resp = _login(c, location)
    assert resp.headers["Location"] == "/pallet/bulk-result?ids=1,2"
    c2 = flask_app.test_client()
    resp = _login(c2, "/login?next=//evil.example.com/x%3Fids%3D1")
    assert "evil" not in resp.headers["Location"]


def test_path_without_query_has_no_trailing_question_mark(client):
    nxt = parse_qs(urlsplit(client.get("/docs?").headers["Location"]).query)["next"][0]
    assert nxt == "/docs"


# ============================================================ FORMS-6


INVOICE_ITEM = {"material_code": "M-1", "qty": "2", "unit_price": "3", "net_weight": "1"}


@pytest.mark.parametrize("url,doc_type", [("/invoice-br/new", "invoice_br"),
                                          ("/invoice-no/new", "invoice_no"),
                                          ("/invoice-dubai/new", "invoice_dubai")])
def test_invoice_without_consignee_is_not_issued(admin_client, db_module, url, doc_type):
    resp = post_with_csrf(admin_client, url, {
        "consignee_name": "   ", "invoice_number": "REQ-1",
        "items_json": json.dumps([INVOICE_ITEM])}, csrf_source_url=url)
    assert resp.status_code == 302 and resp.headers["Location"].startswith(url + "?restore=")
    assert _doc_count(db_module, doc_type) == 0
    body = admin_client.get(resp.headers["Location"]).get_data(as_text=True)
    assert "НЕ е издаден" in body and "„Получател“" in body
    assert "REQ-1" in body, "въведеното е запазено"


# ============================================================ FORMS-7 / FORMS-13


def test_cmr_edit_with_blank_consignee_is_not_saved(admin_client, db_module):
    doc_id = issue_cmr(admin_client, consignee_name="Оригинален Получател")
    resp = post_with_csrf(admin_client, "/doc/%d/edit" % doc_id,
                          {"consignee_name": "   ", "sender_name": "Нов изпращач"},
                          csrf_source_url="/doc/%d/edit" % doc_id)
    assert resp.status_code == 302
    assert resp.headers["Location"].startswith("/doc/%d/edit?restore=" % doc_id)
    assert _doc_data(db_module, doc_id)["consignee_name"] == "Оригинален Получател"
    body = admin_client.get(resp.headers["Location"]).get_data(as_text=True)
    assert "Промените НЕ са записани" in body and "Нов изпращач" in body


def test_invoice_edit_with_blank_consignee_is_not_saved(admin_client, db_module):
    doc_id = _issue(admin_client, "/invoice-no/new", {
        "consignee_name": "ABB Норвегия", "invoice_number": "ED-1",
        "items_json": json.dumps([INVOICE_ITEM])})
    resp = post_with_csrf(admin_client, "/doc/%d/edit" % doc_id, {
        "consignee_name": "", "invoice_number": "ED-1",
        "items_json": json.dumps([INVOICE_ITEM])}, csrf_source_url="/doc/%d/edit" % doc_id)
    assert resp.headers["Location"].startswith("/doc/%d/edit?restore=" % doc_id)
    assert _doc_data(db_module, doc_id)["consignee_name"] == "ABB Норвегия"


def test_number_reuse_on_edit_names_the_edit_button(admin_client):
    _issue(admin_client, "/invoice-no/new", {"consignee_name": "А", "invoice_number": "SAME-1",
                                             "items_json": json.dumps([INVOICE_ITEM])})
    doc_id = _issue(admin_client, "/invoice-br/new", {
        "consignee_name": "Б", "invoice_number": "BR-1", "items_json": json.dumps([INVOICE_ITEM])})
    resp = post_with_csrf(admin_client, "/doc/%d/edit" % doc_id, {
        "consignee_name": "Б", "invoice_number": "SAME-1",
        "items_json": json.dumps([INVOICE_ITEM])}, csrf_source_url="/doc/%d/edit" % doc_id)
    body = html.unescape(admin_client.get(resp.headers["Location"]).get_data(as_text=True))
    assert "натиснете „Запази промените“ още веднъж" in body


# ============================================================ FORMS-10 / дата


def _patch_legacy(monkeypatch, report):
    import routes_settings
    monkeypatch.setattr(routes_settings.legacy_migration, "legacy_names_in_use",
                        lambda *a: True)
    monkeypatch.setattr(routes_settings.updater, "can_complete_rename", lambda: False)
    monkeypatch.setattr(routes_settings.legacy_migration, "last_result", lambda: report)


def test_legacy_names_text_tells_the_real_reason_without_a_missing_button(admin_client,
                                                                          monkeypatch):
    _patch_legacy(monkeypatch, {"status": "kept", "reason": "running", "kind": "in_place"})
    body = html.unescape(admin_client.get("/my-settings").get_data(as_text=True))
    assert "мрежова или споделена папка" not in body
    assert "старата версия на програмата още работеше" in body
    assert "опитано отново при следващото стартиране" in body
    # Без изричен db_path към старото име бутонът го няма — и препратката.
    assert "може да бъде преименувана от раздела" not in body


def test_legacy_names_text_without_a_known_reason_is_neutral(admin_client, monkeypatch):
    _patch_legacy(monkeypatch, None)
    body = html.unescape(admin_client.get("/my-settings").get_data(as_text=True))
    assert "запазват старите си имена (pacho_…) — програмата работи с тях нормално" in body
    assert "мрежова или споделена папка" not in body


def test_instances_last_seen_is_shown_as_a_bulgarian_date(admin_client, db_module):
    con = db_module.get_db()
    con.execute("INSERT INTO app_instances (machine, version, exe_name, last_seen)"
                " VALUES ('OFFICE-PC', '3.81.0', 'PHLogistics.exe', '2026-10-13 15:07:20')")
    con.commit()
    con.close()
    body = admin_client.get("/my-settings").get_data(as_text=True)
    assert "13.10.2026 15:07" in body and "2026-10-13 15:07:20" not in body


# ============================================================ FORMS-11 / FORMS-12 (източник)


def test_single_added_pallet_row_uses_the_singular():
    src = read_source("static", "app.js")
    assert 'tfp("pallet_rows_added_one", "Добавен {count} ред от палетна карта: {list}."' in src
    base = read_source("templates", "base.html")
    assert "'pallet_rows_added_one': _('Добавен {count} ред от палетна карта: {list}.')" in base


def test_bulk_review_sentence_is_one_msgid():
    src = read_source("templates", "pallet_bulk_review.html")
    assert "_('“, „2 от')" not in src
    assert "„1 от %(total)s“, „2 от %(total)s“" in src


# ============================================================ „Копирай като нов“


def test_copy_of_a_packing_list_drops_the_invoice_and_order_numbers(admin_client):
    doc_id = _issue(admin_client, "/packing/new", {
        "receiver_name": "Клиент Копие", "receiver_city": "Анкара", "invoice_no": "0000012955",
        "order_no": "PO-4711", "items_json": json.dumps([{"description": "Кашон", "qty": "1"}])})
    resp = admin_client.get("/doc/%d/copy" % doc_id)
    body = html.unescape(admin_client.get(resp.headers["Location"]).get_data(as_text=True))
    m = re.search(r"data-edit='([^']*)'", body)
    prefill = json.loads(m.group(1))
    assert "invoice_no" not in prefill and "order_no" not in prefill
    assert prefill["receiver_name"] == "Клиент Копие" and prefill["receiver_city"] == "Анкара"
    assert "„Фактура №“, „Поръчка №“" in body


@pytest.mark.parametrize("url,fields,dropped", [
    ("/pallet/new", {"client_name": "К", "ref_cmr": "CMR-1", "pallet_no": "1 от 3",
                     "items_json": json.dumps([{"order_no": "1", "qty": "1"}])},
     ("ref_cmr", "pallet_no")),
    ("/cmr/new", {"consignee_name": "К", "attached_docs": "Invoice 12955"}, ("attached_docs",)),
    ("/export-it/new", {"receiver_name": "К", "invoice_no": "INV-9", "ref_cmr": "CMR-9",
                        "items_json": json.dumps([{"description": "x", "qty": "1"}])},
     ("invoice_no", "ref_cmr")),
])
def test_copy_drops_shipment_references(admin_client, url, fields, dropped):
    doc_id = _issue(admin_client, url, fields)
    resp = admin_client.get("/doc/%d/copy" % doc_id)
    body = html.unescape(admin_client.get(resp.headers["Location"]).get_data(as_text=True))
    prefill = json.loads(re.search(r"data-edit='([^']*)'", body).group(1))
    for key in dropped:
        assert key not in prefill, key
    assert "К" in json.dumps(prefill, ensure_ascii=False)


# ============================================================ диапазон от палетни карти


def _pull(client, url, code):
    return post_with_csrf(client, url, {"code": code}, csrf_source_url="/packing/new").get_json()


@pytest.mark.parametrize("url", ["/packing/pull-pallet", "/invoice/pull-pallet"])
def test_missing_cards_in_a_range_give_one_message(admin_client, url):
    data = _pull(admin_client, url, "6-9")
    assert data["ok"] is False
    assert data["error"].count("Няма палетн") == 1, data["error"]
    assert "6, 7, 8, 9" in data["error"]


@pytest.mark.parametrize("url", ["/packing/pull-pallet", "/invoice/pull-pallet"])
def test_partly_found_range_lists_the_missing_and_returns_them(admin_client, url):
    _issue(admin_client, "/pallet/new", {
        "client_name": "К", "items_format": "orders",
        "items_json": json.dumps([{"order_no": "4700", "pos": "10", "reference": "R",
                                   "qty": "5"}])})
    data = _pull(admin_client, url, "1-3")
    assert data["ok"] is True
    assert data["missing"] == ["2", "3"]
    assert len(data["errors"]) == 1 and "2, 3" in data["errors"][0]


@pytest.mark.parametrize("url", ["/packing/pull-pallet", "/invoice/pull-pallet"])
def test_range_over_the_limit_explains_the_limit(admin_client, url):
    data = _pull(admin_client, url, "1-500")
    assert data["ok"] is False
    assert "„1-500“" in data["error"] and "200" in data["error"]
    assert "Няма палетна карта" not in data["error"]


def test_app_js_keeps_the_missing_cards_in_the_input():
    src = read_source("static", "app.js")
    assert src.count('input.value = (data.missing || []).join(", ");') == 2


# ============================================================ „Общо“ в износа на документите


def test_documents_export_without_type_filter_totals_per_type(admin_client):
    _issue(admin_client, "/cmr/new", {"consignee_name": "А", "weight": "1000", "packages": "4"})
    _issue(admin_client, "/packing/new", {"receiver_name": "А", "total_gross": "1000",
                                          "total_packages": "4",
                                          "items_json": json.dumps([{"description": "x"}])})
    ws = load_workbook(io.BytesIO(admin_client.get("/docs/export.xlsx").data)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    labels = [r[0] for r in rows]
    assert "Общо" not in labels, "сборът на бруто от ЧМР + лист брои стоката два пъти"
    by_label = {r[0]: r for r in rows if isinstance(r[0], str)}
    assert by_label["Общо — ЧМР товарителница"][7] == 1000 and by_label["Общо — ЧМР товарителница"][2] == 1
    assert by_label["Общо — Опаковъчен лист"][7] == 1000
    assert by_label["Общо документи"][2] == 2 and by_label["Общо документи"][7] is None
    assert any(isinstance(x, str) and "поотделно за всеки вид" in x for x in labels)
    ws = load_workbook(io.BytesIO(admin_client.get("/docs/export.xlsx?type=cmr").data)).active
    total = [c.value for c in list(ws.iter_rows())[-1]]
    assert total[0] == "Общо" and total[7] == 1000


# ============================================================ Excel износ на фактурите


def test_invoices_export_follows_the_filters_without_weight(admin_client):
    _issue(admin_client, "/invoice-br/new", {
        "consignee_name": "ABB Бразилия", "invoice_number": "BR-77",
        "items_json": json.dumps([dict(INVOICE_ITEM, po_no="PO-1", qty="2", unit_price="10.5")])})
    _issue(admin_client, "/invoice-no/new", {
        "consignee_name": "ABB Норвегия", "invoice_number": "NO-77",
        "items_json": json.dumps([{"description": "x", "qty": "1", "unit_price": "4"}])})
    resp = admin_client.get("/invoices/export.xlsx")
    assert resp.status_code == 200 and resp.mimetype.endswith("spreadsheetml.sheet")
    ws = load_workbook(io.BytesIO(resp.data)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    header = rows[0]
    assert not [h for h in header if "тегло" in str(h).lower() or "weight" in str(h).lower()]
    money = header.index("Сума, EUR")
    body = {r[2]: r for r in rows[1:-1]}
    assert set(body) == {"BR-77", "NO-77"}
    assert body["BR-77"][money] == 21 and body["BR-77"][5] == "PO-1"
    assert "€" in ws.cell(row=2, column=money + 1).number_format
    assert rows[-1][0] == "Общо" and rows[-1][money] == 25 and rows[-1][2] == 2
    ws = load_workbook(io.BytesIO(
        admin_client.get("/invoices/export.xlsx?type=invoice_no").data)).active
    assert [r[2].value for r in list(ws.iter_rows())[1:-1]] == ["NO-77"]
    ws = load_workbook(io.BytesIO(admin_client.get("/invoices/export.xlsx?q=Бразилия").data)).active
    assert [r[2].value for r in list(ws.iter_rows())[1:-1]] == ["BR-77"]


def test_invoices_page_has_the_excel_button_and_export_needs_login(admin_client, client):
    body = admin_client.get("/invoices").get_data(as_text=True)
    m = re.search(r'<button[^>]*formaction="/invoices/export\.xlsx"[^>]*>', body)
    assert m and 'formmethod="get"' in m.group(0)
    assert client.get("/invoices/export.xlsx").status_code == 302


# ============================================================ обикновени форми при изтекла сесия


def test_new_client_form_is_restored_after_login(flask_app, admin_client, db_module):
    fresh, resp = _expired_post(flask_app, admin_client, "/clients/new", "/clients/new", {
        "name": "Спасен Клиент ООД", "city": "Русе",
        "unload_points_json": json.dumps([{"name": "Склад 2", "address": "ул. 1"}])})
    assert resp.status_code == 302 and "rescue=" in resp.headers["Location"], \
        "без поправката: 400 и въведеното се губи"
    resp = _login(fresh, resp.headers["Location"])
    assert resp.headers["Location"].startswith("/clients/new?restore_form=")
    body = html.unescape(fresh.get(resp.headers["Location"]).get_data(as_text=True))
    assert "Спасен Клиент ООД" in body and "Русе" in body and "Склад 2" in body
    con = db_module.get_db()
    assert con.execute("SELECT COUNT(*) FROM clients").fetchone()[0] == 0
    con.close()


def test_restored_client_edit_keeps_the_submitted_originals(flask_app, admin_client, db_module):
    con = db_module.get_db()
    cid = con.execute("INSERT INTO clients (name, city) VALUES ('Стар', 'София')").lastrowid
    con.commit()
    con.close()
    url = "/clients/%d/edit" % cid
    fresh, resp = _expired_post(flask_app, admin_client, url, url, {
        "name": "Стар", "city": "Варна", "orig_name": "Стар", "orig_city": "София"})
    resp = _login(fresh, resp.headers["Location"])
    body = fresh.get(resp.headers["Location"]).get_data(as_text=True)
    assert 'value="Варна"' in body and 'name="orig_city" value="София"' in body


def test_company_settings_are_restored_and_then_saved(flask_app, admin_client, db_module):
    import db as dbm
    fresh, resp = _expired_post(flask_app, admin_client, "/settings", "/settings", {
        "sender_name": "Фирма", "sender_phone": "+359 888 111", "orig_sender_phone": "",
        "orig_sender_name": "Фирма"})
    resp = _login(fresh, resp.headers["Location"])
    assert resp.headers["Location"].startswith("/settings?restore_form=")
    body = fresh.get(resp.headers["Location"]).get_data(as_text=True)
    assert 'value="+359 888 111"' in body
    assert 'name="orig_sender_phone" value=""' in body, \
        "оригиналът е ПОДАДЕНИЯТ — иначе повторното „Запази“ не записва нищо"
    token = get_csrf_token(fresh, "/settings")
    form = {m.group(1): html.unescape(m.group(2)) for m in
            re.finditer(r'name="((?:orig_)?sender_[a-z_]+)" value="([^"]*)"', body)}
    form["csrf_token"] = token
    fresh.post("/settings", data=form)
    con = dbm.get_db()
    assert dbm.get_settings(con).get("sender_phone") == "+359 888 111"
    con.close()


def test_material_and_invoice_client_forms_are_restored(flask_app, admin_client):
    fresh, resp = _expired_post(flask_app, admin_client, "/materials", "/materials/save", {
        "original_code": "", "code": "MAT-RESCUE", "description": "Спасен", "net_weight": "0,25"})
    resp = _login(fresh, resp.headers["Location"])
    assert resp.headers["Location"].startswith("/materials?restore_form=")
    body = fresh.get(resp.headers["Location"]).get_data(as_text=True)
    assert 'value="MAT-RESCUE"' in body and 'value="0,25"' in body

    fresh, resp = _expired_post(flask_app, admin_client, "/invoices/clients/new",
                                "/invoices/clients/new", {"name": "ABB Спасен",
                                                          "delivery_name": "ABB Ltda"})
    resp = _login(fresh, resp.headers["Location"])
    body = fresh.get(resp.headers["Location"]).get_data(as_text=True)
    assert "ABB Спасен" in body and "ABB Ltda" in body


def test_unknown_form_rescue_token_is_explained(admin_client):
    body = admin_client.get("/settings?restore_form=does-not-exist").get_data(as_text=True)
    assert "не можа да бъде възстановена" in body
