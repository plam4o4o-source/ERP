# -*- coding: utf-8 -*-
"""Одит (15.10.2026, REPORTS): „Справки“ — месечен отчет по клиент
(routes_reports.py, templates/reports.html): правилата за пратки/бруто без
двойно броене, сумите на фактурите в евро, обединяването на клиентите по
регистър/интервали, групирането, правата, Excel износът, линковете към
списъците и плана на заявките (индексът, без да се чете `data`)."""
import io
import json
import random
import re
import sqlite3
from datetime import date
from urllib.parse import parse_qs, urlsplit

import pytest

_SEQ = {"n": 0}


def _add(db_module, doc_type, data, created_at, raw=False):
    """Документ направо в базата (тригерите попълват client_name)."""
    _SEQ["n"] += 1
    n = _SEQ["n"]
    con = sqlite3.connect(db_module.DB_PATH) if raw else db_module.get_db()
    cur = con.execute(
        "INSERT INTO documents (doc_type, number, year, seq, barcode, public_token, data,"
        " created_by, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)",
        (doc_type, "R%04d" % n, int(created_at[:4]), n, "REP-%06d" % n, "rep-%d" % n,
         data if isinstance(data, str) else json.dumps(data, ensure_ascii=False), created_at))
    con.commit()
    doc_id = cur.lastrowid
    con.close()
    return doc_id


class _Recording(sqlite3.Connection):
    log = None

    def execute(self, sql, params=()):
        if _Recording.log is not None:
            _Recording.log.append(sql)
        return super().execute(sql, params)


def _invoice(name, *rows):
    return {"consignee_name": name, "invoice_number": "INV-%s" % name,
            "items": [{"qty": q, "unit_price": p} for q, p in rows]}


@pytest.fixture
def august(db_module):
    """Август 2026: клиент А (три изписвания) — 2 ЧМР с тегло, опаковъчни
    листове и палетни карти със СЪЩАТА стока, фактура; клиент Б — само
    опаковъчни листове; шум извън периода и видове, които не се броят."""
    d = "2026-08-1%d 10:00:00"
    ids = {}
    ids["cmr_a1"] = _add(db_module, "cmr", {"consignee_name": "Клиент А ЕООД", "weight": "1000"}, d % 1)
    ids["cmr_a2"] = _add(db_module, "cmr", {"consignee_name": "клиент а еоод", "weight": "234,5"}, d % 2)
    _add(db_module, "packing", {"receiver_name": "КЛИЕНТ А ЕООД", "total_gross": "1234.5"}, d % 1)
    for i in range(5):
        # едната карта — с двоен интервал в името (същият клиент)
        _add(db_module, "pallet", {"client_name": "Клиент  А ЕООД" if i == 0 else "Клиент А ЕООД",
                                   "gross": "250"}, d % 3)
    ids["inv_a"] = _add(db_module, "invoice_br", _invoice("Клиент А ЕООД", ("2", "10.50"), ("1,5", "3.333")),
                        d % 4)
    _add(db_module, "invoice_no", _invoice("Клиент А ЕООД", ("3", "100")), d % 5)
    # Б: няма ЧМР → пратките и брутото са от опаковъчните листове.
    _add(db_module, "packing", {"receiver_name": "Б ООД", "total_gross": "300"}, d % 1)
    _add(db_module, "packing", {"receiver_name": "Б ООД", "total_gross": "200.5"}, d % 2)
    _add(db_module, "pallet", {"client_name": "Б ООД"}, d % 2)
    # Не влизат: други видове документи и други месеци.
    _add(db_module, "waybill", {"consignee_name": "Клиент А ЕООД"}, d % 3)
    _add(db_module, "dualuse", {"dest_name": "Клиент А ЕООД"}, d % 3)
    _add(db_module, "cmr", {"consignee_name": "Клиент А ЕООД", "weight": "9999"}, "2026-07-31 23:59:59")
    _add(db_module, "cmr", {"consignee_name": "Клиент А ЕООД", "weight": "9999"}, "2026-09-01 00:00:00")
    # Септември — за групирането по месец.
    _add(db_module, "cmr", {"consignee_name": "Клиент А ЕООД", "weight": "500"}, "2026-09-15 08:00:00")
    return ids


def _cells(html):
    """Редовете на таблицата като списъци от текстове на клетките."""
    table = html.split('class="list list-cards report-table"', 1)[1].split("</table>", 1)[0]
    rows = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", table, re.S)[1:]:
        cells = re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)
        rows.append([re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", c)).replace("\xa0", " ").strip()
                     for c in cells])
    return rows


def test_august_rules_without_double_counting(admin_client, august):
    resp = admin_client.get("/reports?from=2026-08&to=2026-08&group=client")
    assert resp.status_code == 200
    rows = _cells(resp.get_data(as_text=True))
    assert rows == [
        # Клиент, Пратки, Палети, Бруто, Фактури, Сума
        ["Б ООД", "2 ОЛ", "1", "500.5 ОЛ", "0", "—"],
        ["Клиент А ЕООД", "2", "5", "1 234.5", "2", "326.00 €"],
        ["Общо", "4", "6", "1 735", "2", "326.00 €"],
    ], rows


def test_invoice_sum_is_the_one_printed_on_the_invoice(db_module, august):
    import appcore
    import routes_reports
    con = db_module.get_db()
    raw = con.execute("SELECT data FROM documents WHERE id = ?", (august["inv_a"],)).fetchone()[0]
    con.close()
    items = json.loads(raw)["items"]
    assert appcore.invoice_totals(items)["price"] == "26.00"  # 21.00 + round(4.9995) = 5.00
    assert routes_reports.invoice_cents(raw) == 2600


def test_invoice_cents_matches_invoice_totals_for_odd_input():
    """Бързият път и резервният (appcore.invoice_row_total) дават ТОЧНО
    сумата под таблицата на фактурата — вкл. запетаи, интервали, отрицателни,
    текст, празни и нечислови редове."""
    import appcore
    import routes_reports
    rnd = random.Random(382)
    pool = ["1", "2", "10", "0", "007", "1,5", "2.25", " 3 ", "1 000", "-1", "abc", "", None, 5, 2.5,
            "1.234,5", "0.005", "12.345", "99999999999999", "3,1", "1e3", "١٢"]
    for _ in range(400):
        items = [{"qty": rnd.choice(pool), "unit_price": rnd.choice(pool)} for _ in range(rnd.randint(0, 6))]
        if rnd.random() < 0.1:
            items.append("развален ред")
        expected = appcore.invoice_totals(items)["price"]
        got = routes_reports.invoice_cents(json.dumps({"items": items}))
        assert (None if got is None else "%d.%02d" % divmod(got, 100)) == (expected or None), items
    assert routes_reports.invoice_cents("{не е json") is None
    assert routes_reports.invoice_cents('{"items": "низ"}') is None


def test_clients_are_merged_by_case_and_whitespace(admin_client, august):
    body = admin_client.get("/reports?from=2026-08&to=2026-08").get_data(as_text=True)
    names = [r[0] for r in _cells(body)]
    variants = [n for n in names if " ".join(n.split()).lower() == "клиент а еоод"]
    assert variants == ["Клиент А ЕООД"], names  # най-честото изписване


def test_grouping_by_month_and_client_month(admin_client, august):
    by_month = _cells(admin_client.get("/reports?from=2026-08&to=2026-09&group=month")
                      .get_data(as_text=True))
    assert [r[0] for r in by_month] == ["август 2026", "септември 2026", "Общо"], by_month
    # август: А 2 ЧМР + Б 2 опаковъчни листа; септември: 2 ЧМР (едното — нощта на 1-ви)
    assert [r[1] for r in by_month] == ["4 ОЛ", "2", "6"]
    assert by_month[1][3] == "10 499"
    cm = _cells(admin_client.get("/reports?from=2026-08&to=2026-09&group=client_month")
                .get_data(as_text=True))
    assert [r[:3] for r in cm] == [["Б ООД", "август 2026", "2 ОЛ"],
                                   ["Клиент А ЕООД", "август 2026", "2"],
                                   ["Клиент А ЕООД", "септември 2026", "2"], ["Общо", "6", "6"]], cm


def test_empty_months_are_shown_when_grouping_by_month(admin_client, august):
    rows = _cells(admin_client.get("/reports?from=2026-06&to=2026-09&group=month").get_data(as_text=True))
    assert [r[0] for r in rows][:4] == ["юни 2026", "юли 2026", "август 2026", "септември 2026"]
    assert rows[0][1:] == ["0", "0", "—", "0", "—"]


def test_huge_period_does_not_build_empty_months(admin_client, august):
    resp = admin_client.get("/reports?from=1900-01&to=9998-12&group=month")
    assert resp.status_code == 200
    assert len(_cells(resp.get_data(as_text=True))) == 4  # юли, август, септември + „Общо“


def test_client_filter_one_client(admin_client, august):
    rows = _cells(admin_client.get("/reports", query_string={
        "from": "2026-08", "to": "2026-08", "client": "  КЛИЕНТ  а еоод "}).get_data(as_text=True))
    assert [r[0] for r in rows] == ["Клиент А ЕООД", "Общо"]
    assert rows[-1][1:] == ["2", "5", "1 234.5", "2", "326.00 €"]


def test_default_period_is_the_current_month(admin_client, db_module):
    today = date.today()
    _add(db_module, "cmr", {"consignee_name": "Днешен Клиент", "weight": "10"},
         today.strftime("%Y-%m-%d 09:00:00"))
    body = admin_client.get("/reports").get_data(as_text=True)
    assert 'value="%s"' % today.strftime("%Y-%m") in body
    assert [r[0] for r in _cells(body)] == ["Днешен Клиент", "Общо"]
    # Бързите бутони: „Този месец“ е отбелязан.
    assert re.search(r'is-active"\s+aria-current="true"\s+href="[^"]*from=%s' % today.strftime("%Y-%m"), body)


def test_reversed_and_invalid_period_is_tolerated(admin_client, august):
    a = _cells(admin_client.get("/reports?from=2026-09&to=2026-08").get_data(as_text=True))
    b = _cells(admin_client.get("/reports?from=2026-08&to=2026-09").get_data(as_text=True))
    assert a == b
    for bad in ("/reports?from=x&to=13-2026", "/reports?from=2026-13", "/reports?group=evil&client=%00"):
        assert admin_client.get(bad).status_code == 200


def test_empty_state(admin_client):
    body = admin_client.get("/reports?from=2001-01&to=2001-02").get_data(as_text=True)
    assert "Няма издадени документи за избрания период." in body
    assert "report-notes" in body


def test_employee_sees_the_same_report_and_anonymous_is_sent_to_login(admin_client, employee_client,
                                                                     client, august):
    url = "/reports?from=2026-08&to=2026-08"
    assert _cells(employee_client.get(url).get_data(as_text=True)) == \
        _cells(admin_client.get(url).get_data(as_text=True))
    for u in (url, "/reports/export.xlsx"):
        resp = client.get(u)
        assert resp.status_code == 302 and "/login" in resp.headers["Location"]


def test_sidebar_link_next_to_all_documents(admin_client, employee_client):
    for c in (admin_client, employee_client):
        body = c.get("/").get_data(as_text=True)
        nav = body.split("<nav>", 1)[1].split("</nav>", 1)[0]
        assert nav.index('href="/docs"') < nav.index('href="/reports"') < nav.index('href="/invoices"')
    body = admin_client.get("/reports").get_data(as_text=True)
    assert re.search(r'class="nav-item active" href="/reports"', body)


def test_client_row_links_to_the_filtered_document_and_invoice_lists(admin_client, august):
    body = admin_client.get("/reports?from=2026-08&to=2026-08").get_data(as_text=True)
    link = re.search(r'<a href="(/docs\?[^"]+)"[^>]*>Клиент А ЕООД</a>', body).group(1).replace("&amp;", "&")
    qs = parse_qs(urlsplit(link).query)
    # Одит (21.10.2026, APPFIX-5): филтър „client“ (ключът на отчета), не текстово q=.
    assert qs == {"client": ["Клиент А ЕООД"], "group": ["client"], "from": ["2026-08-01"], "to": ["2026-08-31"]}
    docs = admin_client.get(link).get_data(as_text=True)
    assert 'href="/doc/%d"' % august["cmr_a1"] in docs and 'href="/doc/%d"' % august["cmr_a2"] in docs
    inv_link = re.search(r'<a href="(/invoices\?[^"]+)"[^>]*>2</a>', body).group(1).replace("&amp;", "&")
    invoices = admin_client.get(inv_link).get_data(as_text=True)
    assert 'href="/doc/%d"' % august["inv_a"] in invoices


def test_excel_export_is_exactly_the_shown_report(admin_client, august):
    from openpyxl import load_workbook
    url = "/reports/export.xlsx?from=2026-08&to=2026-09&group=client_month"
    resp = admin_client.get(url)
    assert resp.status_code == 200
    assert resp.headers["Content-Disposition"].endswith("report_2026-08_2026-09.xlsx")
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert [c.value for c in ws[1]] == ["Клиент", "Месец", "Пратки", "Палети", "Бруто, кг",
                                        "Фактури", "Сума, EUR", "Забележка"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=5)]
    assert rows[1][0] == "Клиент А ЕООД" and rows[1][1].date() == date(2026, 8, 1)
    assert rows[1][2:] == [2, 5, 1234.5, 2, 326.0, None]
    assert rows[0][0] == "Б ООД" and rows[0][2:] == [2, 1, 500.5, 0, None,
                                                     "Пратки; Бруто, кг — по опаковъчните листове"]
    assert rows[3][:1] == ["Общо"] and rows[3][2:] == [6, 6, 12234, 2, 326.0, None]
    assert ws.cell(row=2, column=2).number_format == "mm.yyyy"
    assert ws.cell(row=3, column=7).number_format == '#,##0.00 "€"'
    assert ws.cell(row=3, column=5).number_format == "#,##0.0##"
    assert ws.cell(row=4, column=5).number_format == "#,##0"
    assert ws.cell(row=5, column=1).font.bold and ws.cell(row=5, column=7).number_format == '#,##0.00 "€"'
    assert ws.page_setup.orientation == "landscape" and ws.freeze_panes == "A2"
    notes = [ws.cell(row=r, column=1).value for r in range(6, ws.max_row + 1)]
    assert any(n and n.startswith("Период: 2026-08 – 2026-09") for n in notes)
    assert any(n and n.startswith("Пратки = брой ЧМР") for n in notes)
    # Същите числа като на страницата.
    page = _cells(admin_client.get(url.replace("/export.xlsx", "")).get_data(as_text=True))
    assert page[-1][1:] == ["6", "6", "12 234", "2", "326.00 €"]


def test_excel_export_by_month_has_no_client_column(admin_client, august):
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(admin_client.get(
        "/reports/export.xlsx?from=2026-08&to=2026-08&group=month").data)).active
    assert [c.value for c in ws[1]][:2] == ["Месец", "Пратки"]
    assert ws.cell(row=3, column=1).value == "Общо"


def test_counts_come_from_the_covering_index(db_module, flask_app, august):
    """Броят по вид не чете `data` (покриващият индекс от search_index._INDEXES);
    без унарния плюс (+d.doc_type) планерът избираше индекса по вид документ
    и обхождаше палетните карти от всички години."""
    import routes_reports
    con = sqlite3.connect(db_module.DB_PATH, factory=_Recording)
    _Recording.log = []
    try:
        with flask_app.test_request_context("/reports?from=2026-08"):
            routes_reports.report_cells(con, routes_reports.report_filter())
        executed = list(_Recording.log)
    finally:
        _Recording.log = None
    count_sql = [s for s in executed if "COUNT(*)" in s]
    assert len(count_sql) == 1, executed
    nparams = count_sql[0].count("?")
    plan = " ".join(r[3] for r in con.execute("EXPLAIN QUERY PLAN " + count_sql[0], ["x"] * nparams))
    for sql in executed:  # и другите две заявки минават по created_at
        p = " ".join(r[3] for r in con.execute("EXPLAIN QUERY PLAN " + sql, ["x"] * sql.count("?")))
        assert "idx_documents_created_type_client_number (created_at>? AND created_at<?)" in p, (sql, p)
    con.close()
    assert "COVERING INDEX idx_documents_created_type_client_number" in plan, plan


def test_report_runs_in_one_read_transaction(admin_client, august, monkeypatch):
    """read_snapshot: BEGIN … без запис и без втора връзка (като списъците)."""
    real = sqlite3.connect
    log = []
    _Recording.log = log

    def connect(*a, **kw):
        log.append("CONNECT")
        kw.setdefault("factory", _Recording)
        return real(*a, **kw)

    monkeypatch.setattr(sqlite3, "connect", connect)
    try:
        admin_client.get("/reports?from=2026-08")
        del log[:]
        assert admin_client.get("/reports?from=2026-08").status_code == 200
    finally:
        _Recording.log = None
    assert "BEGIN" in log, log
    inside = log[log.index("BEGIN") + 1:]
    assert inside and "CONNECT" not in inside
    assert not any(re.match(r"\s*(INSERT|UPDATE|DELETE|CREATE|DROP|ALTER)\b", s, re.I) for s in inside)


def test_export_is_audited(admin_client, august, db_module):
    admin_client.get("/reports/export.xlsx?from=2026-08")
    con = db_module.get_db()
    row = con.execute("SELECT action, detail FROM audit_log ORDER BY id DESC LIMIT 1").fetchone()
    con.close()
    assert row[0] == "износ на справката в Excel" and row[1].startswith("2026-08")


def test_documents_without_client_are_last_and_not_linked(admin_client, db_module, august):
    _add(db_module, "cmr", {"consignee_name": "   ", "weight": "7"}, "2026-08-20 10:00:00")
    body = admin_client.get("/reports?from=2026-08&to=2026-08").get_data(as_text=True)
    rows = _cells(body)
    assert rows[-2][:2] == ["(без посочен клиент)", "1"] and rows[-1][0] == "Общо"
    assert not re.search(r"<a [^>]*>\(без посочен клиент\)</a>", body)
