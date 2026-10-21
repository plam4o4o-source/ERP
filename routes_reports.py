# -*- coding: utf-8 -*-
"""„Справки“ — месечен отчет по клиент (одит 15.10.2026, REPORTS).

Заявка: „справка по клиент и месец — брой пратки, палети, бруто тегло,
сума по фактурите (EUR), с износ в Excel; месечен отчет с един клик“.

Правилата за броене (същите текстове стоят като бележки под таблицата и в
Excel файла — виж _notes):
  * Клетката на отчета е „клиент × месец“; групирането по клиент или по
    месец е просто сбор от тези клетки.
  * Пратки = брой ЧМР в клетката; ако клиентът няма нито едно ЧМР в този
    месец — брой опаковъчни листове.
  * Бруто = сбор от кутия 11 на ЧМР; ако в клетката няма ЧМР с попълнено
    тегло — сбор от „Общо бруто“ на опаковъчните листове. Теглата на
    палетните карти и на фактурите НЕ се събират: ЧМР, опаковъчният лист и
    палетните карти на една пратка описват една и съща стока (иначе
    теглото се брои два-три пъти — виж и бележката в documents_export_xlsx).
  * Палети = брой палетни карти; Фактури/Сума EUR — от трите вида фактури
    (сумата е ТОЧНО тази под таблицата на фактурата: appcore.invoice_totals).
  * Месецът е по датата на издаване в програмата (documents.created_at) —
    същото поле, по което филтрират /docs и /invoices, за да съвпадат
    броевете с линковете от отчета.

Права: като списъците „Всички документи“ и „Издадени фактури“ — всеки
влязъл потребител (служител или администратор) вижда всички документи и
фактури (виж shipment._may_view); таблото не показва фактури по изрична
заявка, но тази справка е отделен екран и ги включва.

Регистрира се от routes_documents.register (като shipment.py), за да не се
пипат app.py и списъците с модули в тестовете/сервизните скриптове."""
import hashlib
import io
import re
from collections import Counter
from datetime import date
from urllib.parse import quote

from flask import render_template, request, send_file, url_for
from flask_babel import gettext as _

import applog
import db
from appcore import get_db, invoice_row_total, login_required, read_snapshot, safe_json_data, _parse_decimal

#: Видовете документи, които отчетът чете (останалите не влизат в колоните).
SHIPMENT_TYPES = ("cmr", "packing", "pallet")
REPORT_DOC_TYPES = SHIPMENT_TYPES + tuple(db.INVOICE_DOC_TYPES)
#: Групиране: по клиент, по месец, клиент × месец.
GROUPS = ("client", "month", "client_month")
#: При групиране по месец празните месеци се показват (нула) само до този
#: брой месеци в периода — иначе адрес с период от 1900 до 9999 г. би
#: построил 96 000 празни реда.
FILL_EMPTY_MONTHS_MAX = 36
_MONTH_RE = re.compile(r"(\d{4})-(\d{1,2})(?:-\d{1,2})?\Z")
#: Бърз път за реда на фактура (цяло количество × цена с до 2 знака) —
#: резултатът е ТОЧНО същият като invoice_row_total (произведението е цяло
#: число центове, закръгляне няма); всичко друго минава през invoice_row_total.
_PRICE_RE = re.compile(r"(\d{1,15})(?:[.,](\d{1,2}))?\Z")
_EUR_FORMAT = '#,##0.00 "€"'
_COUNT_FORMAT = "#,##0"
_MONTH_FORMAT = "mm.yyyy"


def register(app):
    app.add_url_rule("/reports", "reports", reports)
    app.add_url_rule("/reports/export.xlsx", "reports_export_xlsx", reports_export_xlsx)


# ---------------------------------------------------------------- помощни

def client_key(name):
    """Ключ за сравнение на клиенти: без излишни интервали и без регистър.
    Одит (21.10.2026, APPFIX-4): db.client_match_key — същото сгъване като
    търсенето в документите (db.search_fold: турското „İ“/„ı“, диакритика),
    общо с филтъра „Клиент“ и „Групирай по клиент“ в /docs и /invoices."""
    return db.client_match_key(name)


def _row_cents(item):
    """Цената на ред от фактура в центове (int) или None — като
    appcore.invoice_row_total (количество × единична цена, закръглено)."""
    if not isinstance(item, dict):
        return None
    qty, price = item.get("qty"), item.get("unit_price")
    if type(qty) is str and type(price) is str and qty.isdigit() and qty.isascii() and len(qty) <= 15:
        m = _PRICE_RE.match(price)
        if m:
            whole, frac = m.groups()
            cents = int(whole) * 100 + (int(frac) * (10 if len(frac) == 1 else 1) if frac else 0)
            return int(qty) * cents
    text = invoice_row_total(item)
    return int(text.replace(".", "")) if text else None


#: Кеш „съдържание на фактурата → сума“ (blake2b-128 на JSON-а): разборът
#: на JSON-а е ~95% от времето на справката за година (≈500 фактури по
#: ~45 реда при 20 000 документа: ~90 ms → ~10 ms при повторно отваряне).
#: Ключът е по САМОТО съдържание, затова редакция на фактурата дава нов ключ
#: (няма остаряване); таванът пази паметта (~4 MB).
_INVOICE_CENTS_CACHE = {}
_INVOICE_CENTS_CACHE_MAX = 20000


def invoice_cents_cached(raw):
    if not isinstance(raw, str):
        return invoice_cents(raw)
    key = hashlib.blake2b(raw.encode("utf-8", "surrogatepass"), digest_size=16).digest()
    try:
        return _INVOICE_CENTS_CACHE[key]
    except KeyError:
        pass
    value = invoice_cents(raw)
    if len(_INVOICE_CENTS_CACHE) >= _INVOICE_CENTS_CACHE_MAX:
        _INVOICE_CENTS_CACHE.clear()
    _INVOICE_CENTS_CACHE[key] = value
    return value


def invoice_cents(raw):
    """Сумата на фактура в центове — същото като appcore.invoice_totals(...)
    ["price"] (сбор от закръглените редове); None, ако няма нито един ред с
    цена (тогава и бланката не показва сума). Регистрира се като SQL
    функция, за да агрегира самата заявка (SUM ... GROUP BY), не Python;
    SUM на само NULL е NULL — така „няма сума“ се различава от 0.00 €."""
    total, priced = 0, False
    for it in safe_json_data(raw).get("items") or []:
        cents = _row_cents(it)
        if cents is not None:
            total += cents
            priced = True
    return total if priced else None


def _weight(text):
    """Тегло от свободен текст (строгият разбор на програмата) или None."""
    num = _parse_decimal(text)
    return num if num is not None and num >= 0 else None


def _register_sql_functions(con):
    con.create_function("rep_client_key", 1, client_key, deterministic=True)
    con.create_function("rep_weight", 1, _weight, deterministic=True)
    con.create_function("rep_invoice_cents", 1, invoice_cents_cached, deterministic=True)


def _parse_month(text):
    m = _MONTH_RE.match((text or "").strip())
    if not m:
        return None
    year, month = int(m.group(1)), int(m.group(2))
    if not (1900 <= year <= 9998 and 1 <= month <= 12):
        return None
    return date(year, month, 1)


def _add_months(d, n):
    y, m = divmod(d.year * 12 + d.month - 1 + n, 12)
    return date(y, m + 1, 1)


def _month_end(d):
    """Последният ден на месеца (за полето „До дата“ на /docs)."""
    return date.fromordinal(_add_months(d, 1).toordinal() - 1)


def _month_label(d):
    from routes_dashboard import _month_name  # късно: същите имена като графиката на таблото
    return _month_name(d, "LLLL y")


def _fmt_count(n):
    return "{:,}".format(int(n or 0)).replace(",", " ")


def _fmt_kg(value):
    if value is None:
        return "—"
    text = "%.3f" % round(value, 3)
    whole, frac = text.split(".")
    frac = frac.rstrip("0")
    whole = "{:,}".format(int(whole)).replace(",", " ")
    return whole + ("." + frac if frac else "")


def _fmt_eur(cents):
    if cents is None:
        return "—"
    return "{:,}.{:02d} €".format(cents // 100, cents % 100).replace(",", " ")


# ---------------------------------------------------------------- филтър

def report_filter(args=None, today=None):
    """Филтърът на отчета от адреса — общ за страницата и Excel износа
    (износът = точно показаният отчет)."""
    args = request.args if args is None else args
    this_month = (today or date.today()).replace(day=1)
    start = _parse_month(args.get("from", "")) or this_month
    end = _parse_month(args.get("to", "")) or (start if args.get("from") else this_month)
    if end < start:
        start, end = end, start
    group = args.get("group", "")
    if group not in GROUPS:
        group = "client"
    client = " ".join((args.get("client", "") or "").split())
    year_start = this_month.replace(month=1)
    last_month = _add_months(this_month, -1)
    quick = [
        ("this_month", _("Този месец"), this_month, this_month),
        ("last_month", _("Миналия месец"), last_month, last_month),
        ("this_year", _("Тази година"), year_start, this_month),
    ]
    return {
        "start": start, "end": end, "end_exclusive": _add_months(end, 1),
        "from": start.strftime("%Y-%m"), "to": end.strftime("%Y-%m"),
        "group": group, "client": client, "client_key": client_key(client) if client else "",
        "quick": [{"key": k, "label": label, "from": a.strftime("%Y-%m"), "to": b.strftime("%Y-%m"),
                   "active": (a, b) == (start, end)} for k, label, a, b in quick],
    }


# ---------------------------------------------------------------- данни

def _new_cell():
    return {"names": Counter(), "cmr": 0, "packing": 0, "pallet": 0, "invoices": 0,
            "cmr_gross": 0.0, "cmr_gross_n": 0, "packing_gross": 0.0, "packing_gross_n": 0,
            "cents": 0, "priced": 0}


def report_cells(con, f):
    """{(ключ на клиента, "ГГГГ-ММ"): клетка} — три GROUP BY заявки по
    индекса idx_documents_created_type_client_number (search_index._INDEXES):
      1) броят по вид — САМО от индекса (covering), без да се чете `data`;
      2) брутото — само ЧМР и опаковъчните листове;
      3) сумите — само фактурите (rep_invoice_cents).
    `+d.doc_type` (унарен плюс) казва на планера да НЕ избира индекса по
    вид документ: без ANALYZE статистика той иначе обхожда всички палетни
    карти от всички години (55 ms вместо 0.6 ms при 20 000 документа)."""
    _register_sql_functions(con)
    where = " WHERE d.created_at >= ? AND d.created_at < ? AND +d.doc_type IN (%s)"
    params = [f["start"].isoformat(), f["end_exclusive"].isoformat()]
    client_sql = ""
    if f["client_key"]:
        # Одит (21.10.2026, APPFIX-4): част от името („ABB“, „muller“) също намира.
        client_sql = " AND instr(rep_client_key(d.client_name), ?) > 0"

    def run(select, types, group_by):
        sql = (select + " FROM documents d" + where % ",".join("?" for _ in types)
               + client_sql + " GROUP BY " + group_by)
        extra = [f["client_key"]] if f["client_key"] else []
        return con.execute(sql, params + list(types) + extra).fetchall()  # nosec B608 -- константни части с „?“ плейсхолдъри

    cells = {}

    def cell(name, ym):
        key = (client_key(name), ym)
        c = cells.get(key)
        if c is None:
            c = cells[key] = _new_cell()
        return c

    for name, ym, doc_type, count in run(
            "SELECT d.client_name, substr(d.created_at, 1, 7), d.doc_type, COUNT(*)",
            REPORT_DOC_TYPES, "1, 2, 3"):
        c = cell(name, ym)
        c["names"][name or ""] += count
        if doc_type in db.INVOICE_DOC_TYPES:
            c["invoices"] += count
        else:
            c[doc_type] += count
    weight_sql = ("rep_weight(CASE WHEN json_valid(d.data) THEN json_extract(d.data,"
                  " CASE d.doc_type WHEN 'cmr' THEN '$.weight' ELSE '$.total_gross' END) END)")
    for name, ym, doc_type, total in run(
            "SELECT d.client_name, substr(d.created_at, 1, 7), d.doc_type, SUM(%s)" % weight_sql,
            ("cmr", "packing"), "1, 2, 3"):
        if total is not None:
            c = cell(name, ym)
            c[doc_type + "_gross"] += total
            c[doc_type + "_gross_n"] += 1
    for name, ym, total in run(
            "SELECT d.client_name, substr(d.created_at, 1, 7), SUM(rep_invoice_cents(d.data))",
            db.INVOICE_DOC_TYPES, "1, 2"):
        if total is not None:
            c = cell(name, ym)
            c["cents"] += total
            c["priced"] += 1
    return cells


def _cell_values(c):
    """Правилата от докстринга на модула за една клетка клиент × месец."""
    shipments_from_packing = not c["cmr"] and c["packing"] > 0
    if c["cmr_gross_n"]:
        gross, gross_from_packing = c["cmr_gross"], False
    elif c["packing_gross_n"]:
        gross, gross_from_packing = c["packing_gross"], True
    else:
        gross, gross_from_packing = None, False
    return {
        "shipments": c["packing"] if shipments_from_packing else c["cmr"],
        "shipments_from_packing": shipments_from_packing,
        "gross": gross, "gross_from_packing": gross_from_packing,
        "pallets": c["pallet"], "invoices": c["invoices"],
        "cents": c["cents"] if c["priced"] else None,
    }


def _add(total, v):
    total["shipments"] += v["shipments"]
    total["pallets"] += v["pallets"]
    total["invoices"] += v["invoices"]
    total["shipments_from_packing"] = total["shipments_from_packing"] or v["shipments_from_packing"]
    total["gross_from_packing"] = total["gross_from_packing"] or v["gross_from_packing"]
    if v["gross"] is not None:
        total["gross"] = (total["gross"] or 0.0) + v["gross"]
    if v["cents"] is not None:
        total["cents"] = (total["cents"] or 0) + v["cents"]


def _empty_total():
    return {"shipments": 0, "pallets": 0, "invoices": 0, "gross": None, "cents": None,
            "shipments_from_packing": False, "gross_from_packing": False,
            "client_key": None, "client": None, "month": None}


def _display_name(counter):
    """Най-често срещаното изписване на клиента (при равенство — първото
    по азбучен ред), за да не се показва случаен вариант."""
    if not counter:
        return ""
    return sorted(counter.items(), key=lambda kv: (-kv[1], db._ci_lower(kv[0]), kv[0]))[0][0]


def build_report(con, f):
    """Редовете на отчета според групирането + ред „Общо“."""
    cells = report_cells(con, f)
    names = {}
    for (key, _ym), c in cells.items():
        names.setdefault(key, Counter()).update(c["names"])
    display = {key: _display_name(counter) for key, counter in names.items()}
    group = f["group"]
    rows = {}
    for (key, ym), c in cells.items():
        month = date(int(ym[:4]), int(ym[5:7]), 1)
        row_key = {"client": (key, None), "month": (None, month), "client_month": (key, month)}[group]
        row = rows.get(row_key)
        if row is None:
            row = rows[row_key] = _empty_total()
            row["client_key"], row["month"] = row_key
            row["client"] = display.get(key, "") if row_key[0] is not None else None
        _add(row, _cell_values(c))
    if group == "month":
        span = (f["end"].year - f["start"].year) * 12 + f["end"].month - f["start"].month + 1
        if span <= FILL_EMPTY_MONTHS_MAX:
            for i in range(span):
                month = _add_months(f["start"], i)
                if (None, month) not in rows:
                    row = rows[(None, month)] = _empty_total()
                    row["client_key"], row["month"], row["client"] = None, month, None

    def order(row):
        client = row["client"]
        client_part = ((client == "" if client is not None else False),
                       db._ci_lower(client or ""), client or "")
        return (client_part, row["month"] or date.min)

    ordered = sorted(rows.values(), key=order)
    total = _empty_total()
    bases = (url_for("documents"), url_for("invoices_list"))
    for row in ordered:
        _add(total, row)
        row.update(_row_links(f, row, bases))
    return {"rows": ordered, "total": total}


def _query_url(base, **params):
    """Адрес с параметри (без празните) — като url_for, но без да строи
    маршрута за всеки ред (400 реда × 2 url_for бяха ~25 ms)."""
    parts = ["%s=%s" % (k, quote(str(v), safe="")) for k, v in params.items() if v]
    return base + ("?" + "&".join(parts) if parts else "")


def _row_links(f, row, bases):
    """Линковете на реда към /docs и /invoices със същия клиент и период
    (параметрите на филтрите там: q, from, to, group — виж
    routes_documents._documents_list_filter/routes_invoices._invoices_list_filter)."""
    start, end = (row["month"], row["month"]) if row["month"] else (f["start"], f["end"])
    period = {"from": start.isoformat(), "to": _month_end(end).isoformat()}
    # Одит (21.10.2026, APPFIX-5): филтър „client“ по СЪЩИЯ ключ като отчета
    # (не текстово търсене q=…) — „ЕНЕРГО-ПРО   МРЕЖИ АД“ влиза в реда.
    client = row["client"] or ""
    return {"docs_url": _query_url(bases[0], client=client, group="client" if client else "", **period),
            "invoices_url": _query_url(bases[1], client=client, **period)}


def client_choices(con):
    """Клиентите за полето „Клиент“ (datalist): едно изписване на ключ, по
    индекса idx_documents_client_type (covering — без `data`)."""
    counts = {}
    sql = ("SELECT client_name, COUNT(*) FROM documents WHERE client_name != ''"
           " AND +doc_type IN (%s) GROUP BY client_name" % ",".join("?" for _ in REPORT_DOC_TYPES))  # nosec B608 -- само „?“ плейсхолдъри
    for name, count in con.execute(sql, list(REPORT_DOC_TYPES)):
        counts.setdefault(client_key(name), Counter())[name] += count
    names = [_display_name(c) for c in counts.values()]
    return sorted(names, key=lambda n: (db._ci_lower(n), n))


def _notes():
    return [
        _("Пратки = брой ЧМР. Ако за клиента в даден месец няма нито едно ЧМР, "
          "се броят опаковъчните листове (отбелязано с „ОЛ“)."),
        _("Бруто, кг = сбор от кутия 11 на ЧМР. Ако за клиента в даден месец няма ЧМР "
          "с попълнено тегло, се взема „Общо бруто“ от опаковъчните листове. Теглата "
          "на палетните карти не се добавят — те описват същата стока."),
        _("Палети = брой палетни карти. Сума EUR = сбор от общите суми на фактурите "
          "(Бразилия, Норвегия, Дубай)."),
        _("Месецът е по датата на издаване в програмата; клиентът — по името в "
          "документа, без разлика в главни/малки букви и интервали."),
    ]


# ---------------------------------------------------------------- изгледи

@login_required
@read_snapshot  # само чете — една транзакция (appcore.read_snapshot)
def reports():
    f = report_filter()
    con = get_db()
    report = build_report(con, f)
    for row in report["rows"] + [report["total"]]:
        row["shipments_text"] = _fmt_count(row["shipments"])
        row["pallets_text"] = _fmt_count(row["pallets"])
        row["invoices_text"] = _fmt_count(row["invoices"])
        row["gross_text"] = _fmt_kg(row["gross"])
        row["amount_text"] = _fmt_eur(row["cents"])
        row["month_text"] = _month_label(row["month"]) if row["month"] else ""
    period = (_month_label(f["start"]) if f["start"] == f["end"]
              else "%s – %s" % (_month_label(f["start"]), _month_label(f["end"])))
    return render_template("reports.html", f=f, report=report, period=period,
                           clients=client_choices(con), notes=_notes())


def _headers(group):
    cols = []
    if group != "month":
        cols.append(("client", _("Клиент"), 34))
    if group != "client":
        cols.append(("month", _("Месец"), 12))
    cols += [("shipments", _("Пратки"), 10), ("pallets", _("Палети"), 10),
             ("gross", _("Бруто, кг"), 14), ("invoices", _("Фактури"), 10),
             ("amount", _("Сума, EUR"), 16),
             # „ОЛ“ от страницата: кои стойности са по опаковъчните листове
             ("note", _("Забележка"), 36)]
    return cols


def _kg_format(value):
    return _COUNT_FORMAT if float(value).is_integer() else "#,##0.0##"


@login_required
def reports_export_xlsx():
    """Excel износ на ТОЧНО показания отчет (същият филтър/групиране) — със
    стила на износите на списъците (routes_documents._list_*): числата са
    числа, сумите в евро, ред „Общо“, бележките отдолу, печат A4 легнало."""
    from openpyxl import Workbook

    from routes_documents import (_list_row_styler, _list_sheet_finish, _list_sheet_header,
                                  _list_total_row, _xlsx_append, _xlsx_styles)

    f = report_filter()
    con = get_db()
    report = build_report(con, f)
    cols = _headers(f["group"])
    keys = [k for k, _l, _w in cols]
    st = _xlsx_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = _("Справка")[:31]
    last = _list_sheet_header(ws, [label for _k, label, _w in cols], st)
    style_row = _list_row_styler(ws, st, len(cols))
    no_client = _("(без посочен клиент)")

    def values_of(row):
        out, formats = [], {}
        for i, key in enumerate(keys, start=1):
            if key == "client":
                out.append(row["client"] or no_client)
            elif key == "month":
                out.append(row["month"])
                formats[i] = _MONTH_FORMAT
            elif key == "gross":
                value = round(row["gross"], 3) if row["gross"] is not None else ""
                out.append(value)
                if value != "":
                    formats[i] = _kg_format(value)
            elif key == "note":
                parts = [label for flag, label in ((row["shipments_from_packing"], _("Пратки")),
                                                   (row["gross_from_packing"], _("Бруто, кг"))) if flag]
                out.append("%s — %s" % ("; ".join(parts), _("по опаковъчните листове")) if parts else "")
            elif key == "amount":
                value = row["cents"] / 100 if row["cents"] is not None else ""
                out.append(value)
                if value != "":
                    formats[i] = _EUR_FORMAT
            else:
                out.append(row[key])
                formats[i] = _COUNT_FORMAT
        return out, formats

    for row in report["rows"]:
        values, formats = values_of(row)
        last = _xlsx_append(ws, values)
        style_row(last)
        for c, fmt in formats.items():
            ws.cell(row=last, column=c).number_format = fmt
    values, formats = values_of(report["total"])
    # „Общо“ в първата колона; колоните клиент/месец на реда остават празни.
    values[keys.index("note")] = ""
    for i, key in enumerate(keys, start=1):
        if key in ("client", "month"):
            values[i - 1] = _("Общо") if i == 1 else ""
            formats.pop(i, None)
    last = _list_total_row(ws, st, len(cols), values, {})
    for c, fmt in formats.items():
        ws.cell(row=last, column=c).number_format = fmt
    period = "%s – %s" % (f["from"], f["to"]) if f["from"] != f["to"] else f["from"]
    _xlsx_append(ws, [])
    _xlsx_append(ws, ["%s: %s" % (_("Период"), period)
                      + ("; %s: %s" % (_("Клиент"), f["client"]) if f["client"] else "")])
    for note in _notes():
        _xlsx_append(ws, [note])
    _list_sheet_finish(ws, len(cols), [w for _k, _l, w in cols], last)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    applog.log_audit("износ на справката в Excel", "%s, %d реда" % (period, len(report["rows"])))
    return send_file(buf, as_attachment=True,
                     download_name="report_%s_%s.xlsx" % (f["from"], f["to"]),
                     mimetype="application/vnd.openxmlformats-officedocument"
                              ".spreadsheetml.sheet")
